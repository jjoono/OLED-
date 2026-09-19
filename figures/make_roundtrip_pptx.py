# -*- coding: utf-8 -*-
"""Build Fig.1(a) as a PowerPoint slide made of native, individually editable shapes.

Nothing is grouped and nothing is a picture: every lens is an Oval, every ray is a
straight arrow connector, every star is a PowerPoint star, every squiggle is a
freeform, and all text sits in real text boxes.  Geometry and the radiometry are
imported from make_roundtrip_figure.py, so the .pptx and the .svg cannot drift.
"""
import math, os, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import make_roundtrip_figure as F          # geometry + colours, single source of truth

from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

# ------------------------------------------------------------------ canvas
SLIDE_W, SLIDE_H = Inches(13.333), Inches(2.90)
FIG_W = Inches(13.0)
SCALE = FIG_W / float(F.W)                       # EMU per SVG user unit
FIG_H = F.H * SCALE
OX = (SLIDE_W - FIG_W) / 2.0
OY = (SLIDE_H - FIG_H) / 2.0

def X(u): return int(round(OX + u * SCALE))
def Y(u): return int(round(OY + u * SCALE))
def D(u): return int(round(u * SCALE))
def PTS(u): return Pt(u * SCALE / 12700.0)

C = lambda h: RGBColor.from_string(h.lstrip("#").upper())
GREEN, BLUE, AMBER = C(F.RAY), C(F.LOSS), C(F.EMIT)      # light / loss / emission
INK, MUTED, LEADER = C(F.INK), C(F.MUTED), C(F.LEADER)
SUB_FILL, SUB_LINE = C(F.GLASS_F), C(F.GLASS_L)

prs = Presentation()
prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
slide = prs.slides.add_slide(prs.slide_layouts[6])      # blank
SH = slide.shapes

# ------------------------------------------------------------------ helpers
def _alpha(parent, a):
    """Apply transparency (a = opacity 0..1) to a solidFill under `parent`."""
    if a is None or a >= 0.999:
        return
    sf = parent.find(qn("a:solidFill"))
    if sf is None:
        return
    clr = sf.find(qn("a:srgbClr"))
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(a * 100000))}))

def style(shape, fill=None, line=None, lw=None, op=None, name=None):
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
        _alpha(shape._element.spPr, op)
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        if lw is not None:
            shape.line.width = lw
        _alpha(shape._element.spPr.find(qn("a:ln")), op)
    shape.shadow.inherit = False
    if name:
        shape.name = name
    return shape

def arrow(p0, p1, w, color=GREEN, head=True, op=None, name="ray"):
    cn = SH.add_connector(MSO_CONNECTOR.STRAIGHT, X(p0[0]), Y(p0[1]), X(p1[0]), Y(p1[1]))
    cn.line.color.rgb = color
    cn.line.width = Emu(max(3175, D(w)))            # >= 0.25 pt so PowerPoint keeps it
    ln = cn.line._get_or_add_ln()
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"),
                                 {"type": "triangle", "w": "lg", "len": "med"}))
    _alpha(ln, op)
    cn.shadow.inherit = False
    cn.name = name
    return cn

def freewave(x0, y, dx, amp, periods, color, lw, op, name="loss-wave", head=False):
    n = 48
    pts = [(X(x0 + dx * i / n), Y(y + amp * math.sin(2 * math.pi * periods * i / n)))
           for i in range(n + 1)]
    b = SH.build_freeform(pts[0][0], pts[0][1], scale=1.0)
    b.add_line_segments(pts[1:], close=False)
    s = b.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = color
    s.line.width = Emu(max(3175, D(lw)))
    ln = s.line._get_or_add_ln()
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"),
                                 {"type": "triangle", "w": "med", "len": "med"}))
    _alpha(ln, op)
    s.shadow.inherit = False
    s.name = name
    return s

def star(cx, cy, r, shape_kind, fill, line, op=None, lw=1.3, name="star"):
    s = SH.add_shape(shape_kind, X(cx - r), Y(cy - r), D(2 * r), D(2 * r))
    try:
        s.adjustments[0] = 0.40          # inner radius, matches the SVG stars
    except (IndexError, ValueError):
        pass
    return style(s, fill, line, Emu(max(3175, D(lw))), op, name)

def text(cx_or_x, baseline, s, size, color, align=PP_ALIGN.CENTER,
         bold=False, box_w=600.0, name="label"):
    left = X(cx_or_x - box_w / 2.0) if align == PP_ALIGN.CENTER else (
        X(cx_or_x - box_w) if align == PP_ALIGN.RIGHT else X(cx_or_x))
    h = size * 2.4
    tb = SH.add_textbox(left, Y(baseline - 0.34 * size - h / 2.0), D(box_w), D(h))
    tf = tb.text_frame
    tf.word_wrap = True                 # wrap="none" would override `align`
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = s
    r.font.size = PTS(size)
    r.font.bold = bold
    r.font.name = "Arial"
    r.font.color.rgb = color
    tb.name = name
    return tb

# ------------------------------------------------------------------ one panel
def panel(x0, eta, r_met, t_tco, tag):
    # lenses first: full circles, then the substrate slab is laid over their lower halves
    for k in range(int(F.PW / F.PITCH)):
        cx = x0 + F.PITCH * k + F.LENS_R
        style(SH.add_shape(MSO_SHAPE.OVAL, X(cx - F.LENS_R), Y(F.Y_TOP - F.LENS_R),
                           D(2 * F.LENS_R), D(2 * F.LENS_R)),
              SUB_FILL, SUB_LINE, PTS(F.LW_SLAB), name="%s lens %02d" % (tag, k + 1))
    style(SH.add_shape(MSO_SHAPE.RECTANGLE, X(x0), Y(F.Y_TOP), D(F.PW), D(F.Y_TCO - F.Y_TOP)),
          SUB_FILL, SUB_LINE, PTS(F.LW_SLAB), name="%s glass substrate" % tag)
    for y1, y2, fl, st, nm in ((F.Y_TCO, F.Y_ORG, F.TCO_F, F.TCO_L, "TCO anode"),
                               (F.Y_ORG, F.Y_MET, F.ORG_F, F.ORG_L, "organic layers"),
                               (F.Y_MET, F.Y_BOT, F.MET_F, F.MET_L, "metal cathode")):
        style(SH.add_shape(MSO_SHAPE.RECTANGLE, X(x0), Y(y1), D(F.PW), D(y2 - y1)),
              C(fl), C(st), PTS(F.LW_LAYER), name="%s %s" % (tag, nm))

    # radiometry - identical formulas to the SVG generator
    I, seg, esc, loss = 1.0, [], [], []
    for k in range(4):
        seg.append(I); esc.append(I * eta); I *= (1 - eta)
        if k < 3:
            I *= t_tco; loss.append(I * (1 - r_met)); I *= r_met; I *= t_tco
    wid = [F.W_RAY * v ** 0.75 for v in seg]
    kept = (1.0 - eta) ** 0.75

    tl = 1.0 - t_tco
    tw, ta = 10.0 + 180.0 * tl, 1.0 + 22.0 * tl
    for i, xm in enumerate(F.HIT_METAL):
        for j, sgn in enumerate((-1, 1)):
            freewave(x0 + xm + sgn * 24 - tw / 2, F.Y_TCO + 8, tw, ta, 2, BLUE,
                     0.9 + 9.0 * tl, min(1.0, 0.25 + 6.5 * tl),
                     name="%s TCO absorption %d%s" % (tag, i + 1, "LR"[j]))

    arrow((x0 + F.X_EMIT + 7.1, F.Y_EMIT - 7.1), (x0 + F.HIT_TOP[0], F.Y_TOP), wid[0],
          name="%s ray emit" % tag)
    for k in range(3):
        arrow((x0 + F.HIT_TOP[k] + 9.2, F.Y_TOP + 9.2),
              (x0 + F.HIT_METAL[k] - 6.4, F.Y_MET - 6.4), wid[k] * kept,
              name="%s ray down %d" % (tag, k + 1))
        arrow((x0 + F.HIT_METAL[k] + 6.4, F.Y_MET - 6.4),
              (x0 + F.HIT_TOP[k + 1], F.Y_TOP), wid[k + 1],
              name="%s ray up %d" % (tag, k + 1))
    arrow((x0 + F.HIT_TOP[3] + 9, F.Y_TOP + 9), (x0 + F.PW, F.Y_TOP + F.PW - F.HIT_TOP[3]),
          wid[3] * kept, head=False, op=0.45, name="%s ray continues" % tag)

    for k, xl in enumerate(F.HIT_TOP):
        fw = max(1.0, 0.55 * wid[k])
        for j, (phi, ln_) in enumerate(((-30, 68), (-5, 79), (22, 70))):
            a = math.radians(phi); b = math.radians(phi * 1.45)
            o = (x0 + xl + F.LENS_R * math.sin(a), F.Y_TOP - F.LENS_R * math.cos(a))
            arrow(o, (o[0] + ln_ * math.sin(b), o[1] - ln_ * math.cos(b)), fw,
                  name="%s outcoupled %d-%d" % (tag, k + 1, j + 1))

    for k, xm in enumerate(F.HIT_METAL):
        v = loss[k]
        r = 3.0 + 30.0 * v ** 0.6
        op = 0.5 + 0.5 * min(1.0, v / 0.18)
        for j, sgn in enumerate((-1, 1)):
            freewave(x0 + xm + sgn * (r - 1), F.Y_MET - 6, sgn * (10.0 + 200.0 * v),
                     2.0 + 14 * v, 2.5, BLUE, 1.2 + 6.0 * v ** 0.6, op, head=True,
                     name="%s ohmic loss %d%s" % (tag, k + 1, "LR"[j]))
        star(x0 + xm, F.Y_MET, r, MSO_SHAPE.STAR_10_POINT, C(F.BURST), BLUE, op,
             F.LW_STAR, "%s absorption burst %d" % (tag, k + 1))

    star(x0 + F.X_EMIT, F.Y_EMIT, 8.5, MSO_SHAPE.STAR_8_POINT, AMBER, C(F.EMIT_EDGE),
         None, F.LW_STAR, "%s emitter" % tag)
    style(SH.add_shape(MSO_SHAPE.OVAL, X(x0 + F.X_EMIT - 2.4), Y(F.Y_EMIT - 2.4),
                       D(4.8), D(4.8)), C("FFF8E6"), None, name="%s emitter core" % tag)
    return sum(esc)

e1 = panel(F.PANEL_X[0], .30, .72, .90, "L")
e2 = panel(F.PANEL_X[1], .30, .98, .995, "R")

text(F.PANEL_X[0] + F.PW / 2, 28, "Conventional OLED + outcoupling structure", 21,
     INK, bold=True, box_w=700.0, name="title left")
text(F.PANEL_X[0] + F.PW / 2, 52,
     "large absorption per round trip — the beam dies out within a few passes",
     16, MUTED, name="subtitle left")
text(F.PANEL_X[1] + F.PW / 2, 28, "Designed by the proposed design rule", 21,
     INK, bold=True, box_w=700.0, name="title right")
text(F.PANEL_X[1] + F.PW / 2, 52,
     "absorption suppressed — intensity survives many round trips",
     16, MUTED, name="subtitle right")

for y_lab, y_tip, s in ((150, 152, "Outcoupling structure"), (222, 218, "Glass substrate"),
                        (263, 269, "TCO anode"), (292, 286, "Organic layers"),
                        (318, 315, "Metal cathode")):
    text(181, y_lab + 5, s, 14, MUTED, PP_ALIGN.RIGHT, box_w=220.0, name="key: " + s)
    c = arrow((187, y_lab), (195, y_tip), 0.9, LEADER, head=False,
              name="key leader: " + s)

LG = 412.0
arrow((642, LG), (706, LG), 6.2, name="legend ray")
text(720, LG + 5, "Light ray  (line width ∝ optical power)", 15, MUTED,
     PP_ALIGN.LEFT, box_w=400.0, name="legend ray text")
star(1070, LG - 2, 8.5, MSO_SHAPE.STAR_8_POINT, AMBER, C(F.EMIT_EDGE), None, F.LW_STAR,
     "legend emitter")
text(1088, LG + 5, "Exciton emission", 15, MUTED, PP_ALIGN.LEFT, box_w=300.0,
     name="legend emission text")
freewave(1300, LG - 2, 56, 3.4, 2.5, BLUE, 2.2, 0.95, name="legend loss", head=True)
text(1376, LG + 5, "Absorption loss  (ohmic at the metal, TCO)", 15, MUTED,
     PP_ALIGN.LEFT, box_w=460.0, name="legend loss text")

dest = os.path.join(os.path.dirname(os.path.abspath(__file__)), "outcoupling_roundtrip.pptx")
prs.save(dest)
print("wrote %s\n  %d shapes, none grouped   |  extracted: %.0f%% vs %.0f%%"
      % (dest, len(SH), 100 * e1, 100 * e2))
