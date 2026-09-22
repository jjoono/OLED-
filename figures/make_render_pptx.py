# -*- coding: utf-8 -*-
"""Put each Cycles render on a slide with its layer labels as real text boxes.

The render itself stays a picture.  A volumetric glow and 331 shaded micro-lens
domes have no vector form worth having -- auto-tracing them yields thousands of
meaningless paths, a larger file and a worse picture.  What is worth vectorising
is the text, and that is what this does: the device is rendered with
`--no-labels`, and the labels go on top as native PowerPoint text boxes.

Their positions come from `*_labels.json`, which the render writes by projecting
each label's anchor through the camera that rendered the frame.  So the text
lands where the 3D text would have, and it follows the camera and the layer
thicknesses without anyone re-measuring anything.

    blender -b -P figures/build_emitting_area_blend.py -- --out figures --no-labels
    python3 figures/make_render_pptx.py
"""
import json
import os
import sys

from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = sys.argv[1] if len(sys.argv) > 1 else HERE
OUT = os.path.join(HERE, "emitting_area_render.pptx")

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.35)
INK_DARK, INK_LIGHT = RGBColor(0x1B, 0x20, 0x26), RGBColor(0xF2, 0xF5, 0xF8)

PANELS = [("emitting_area_render_conventional", "Conventional"),
          ("emitting_area_render_designrule", "Design rule")]


def add(prs, stem, title):
    png = os.path.join(SRC, stem + ".png")
    meta = os.path.join(SRC, stem + "_labels.json")
    if not (os.path.exists(png) and os.path.exists(meta)):
        raise SystemExit("missing %s or its labels -- render with --no-labels first" % stem)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    iw, ih = Image.open(png).size
    h = SLIDE_H - 2 * MARGIN                       # fit to height, centred
    w = int(round(h * iw / float(ih)))
    left, top = int((SLIDE_W - w) / 2), int(MARGIN)
    pic = slide.shapes.add_picture(png, left, top, width=w, height=h)
    pic.name = "%s render" % title

    for lab in json.load(open(meta, encoding="utf-8")):
        # "h" is the projected height of Blender's font size, which is the em --
        # not the cap height.  Converting it as a cap height oversizes the text.
        size = Pt(lab["h"] * (h / 12700.0))
        box_h = int(size * 2.2)
        tb = slide.shapes.add_textbox(left + int(lab["x"] * w),
                                      top + int(lab["y"] * h) - box_h // 2,
                                      int(w * 0.62), box_h)
        tf = tb.text_frame
        tf.word_wrap = True                        # wrap="none" would override `align`
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        r = p.add_run()
        r.text = lab["text"]
        r.font.size = size
        r.font.name = "Arial"
        r.font.color.rgb = INK_LIGHT if lab["light"] else INK_DARK
        tb.name = "%s / %s" % (title, lab["text"])
    return slide


prs = Presentation()
prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
for stem, title in PANELS:
    add(prs, stem, title)
prs.save(OUT)
print("saved %s  (%d slides, text editable, render placed as a picture)"
      % (OUT, len(prs.slides._sldIdLst)))
