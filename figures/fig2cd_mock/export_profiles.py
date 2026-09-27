"""Raw data of the depth profiles in fig2cd_mock.png / figS_dbr_mock.png (net Poynting flux vs depth,
spectrum- and Lambertian-weighted, s/p averaged), using the same code and parameters as fig2cd_mock.py."""
import numpy as np, openpyxl
from openpyxl.styles import Font
src = open('fig2cd_mock.py').read(); exec(src[:src.index('plt.rcParams.update')])
B = Font(bold=True)
cases = [('Al_TCOk0.02', stack('Al', 0.02)), ('Ag_TCOk0.02', stack('Ag', 0.02)), ('Ag_TCOk0.005', stack('Ag', 0.005)),
         ('Al_TCOk0.005', stack('Al', 0.005)), ('DBR_air_TCOk0.02', stack('DBR', 0.02, 'air')), ('DBR_glass_TCOk0.02', stack('DBR', 0.02, 'glass'))]
wb = openpyxl.Workbook(); w = wb.active; w.title = 'README'
for i, l in enumerate(['Depth profiles of the net Poynting flux S_z (normalised to the returned flux) behind fig2cd_mock.png panels (d) and figS_dbr_mock.png.',
    'Same code as figures/fig2cd_mock/fig2cd_mock.py: substrate n = 1.77, TCO 50 nm (n = 2.0 + ik), organics 300 nm (n = 1.8), metal 100 nm (textbook n,k), air;',
    'DBR: TCO2 50 nm + 4.5 pairs ZnS/LiF quarter-wave at 540 nm. Weighting: Gaussian EL 530 nm (FWHM 60 nm) x Lambertian cos.sin over 0-89 deg, s/p average.',
    'Depth z = 0 at the substrate/TCO interface; the substrate part (z < 0) is the constant incident net flux = A\' (all light not reflected).',
    'Loss in a layer = S_z at its entrance - S_z at its exit (sheet "summary").',
    'Note: this is the illustrative draft (textbook n,k). The library-constant version (Koenig ITO, McPeak Ag) is sim/poynting/layer_loss_poynting.xlsx.'], 1): w.cell(i, 1, l)
w.column_dimensions['A'].width = 150
sm = wb.create_sheet('summary'); hdr = ['case', "A' total (%)"]; rowsum = []
for name, L in cases:
    zabs, acc = weighted_profile(L); names = [l[0] for l in L]
    s = wb.create_sheet(name); s.cell(1, 1, 'z (nm)').font = B; s.cell(1, 2, 'S_z / returned flux').font = B; s.cell(1, 3, 'layer').font = B
    r = 2
    for j in range(len(zabs)):
        for zz, ss in zip(zabs[j], acc[j]):
            s.cell(r, 1, round(float(zz), 2)); s.cell(r, 2, float(ss)); s.cell(r, 3, names[j]); r += 1
    loss = {}
    for j in range(1, len(acc) - 1):
        k = 'DBR' if names[j] in ('ZnS', 'LiF') else names[j]; loss[k] = loss.get(k, 0) + acc[j][0] - acc[j][-1]
    rowsum.append((name, acc[0][0], loss, acc[-1][0]))
keys = []
for _, _, l, _ in rowsum:
    for k in l:
        if k not in keys: keys.append(k)
for j, h in enumerate(hdr + [f'{k} (%)' for k in keys] + ['leakage to rear (%)'], 1): sm.cell(1, j, h).font = B
for i, (n, A0, l, T) in enumerate(rowsum, 2):
    vals = [n, 100*A0] + [100*l.get(k, 0) for k in keys] + [100*T]
    for j, v in enumerate(vals, 1): sm.cell(i, j, v if j == 1 else round(float(v), 3))
    print(n, round(100*A0, 2), {k: round(100*v, 2) for k, v in l.items()}, 'leak', round(100*T, 2))
wb.save('fig2cd_profiles_rawdata.xlsx')
