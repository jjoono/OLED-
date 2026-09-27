"""Depth profiles of the net Poynting flux with the manuscript optical constants:
Koenig ITO, McPeak Ag, library Al (JO), organics n = 1.8 (420 nm), glass n = 1.5, metal 100 nm, air;
Ir(ppy)2acac spectrum (photon number, 400-700 nm) x Lambertian cos.sin, s/p average."""
import numpy as np, sys, openpyxl
import matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
from openpyxl.styles import Font
src = open('fig2cd_mock.py').read(); exec(src[:src.index('plt.rcParams.update')])
sys.path.insert(0, '/home/user/OLED-/sim/design_rule4'); import fig2d as F
WL = np.arange(400.0, 701.0, 5.0); sel = np.isin(F.LAM, WL)
Ag, Al, ITO = F.AG[sel], F.AL[sel], F.ITO[sel]
ORG, AIR = 1.8 + 0j*WL, 1.0 + 0j*WL
N_SUB = 1.5
EL = np.clip(np.interp(WL, F.LAM, F.GREEN), 0, None)*WL; EL /= EL.sum()
def stk(metal, d_ito):
    M = Ag if metal == 'Ag' else Al
    return [('substrate', N_SUB + 0j*WL, None), ('ITO', ITO, float(d_ito)), ('organics', ORG, 420.0), (metal, M, 100.0), ('air', AIR, None)]
cases = [(f'{m}_ITO{d}', stk(m, d)) for m in ('Al', 'Ag') for d in (50, 150)]
B = Font(bold=True); wb = openpyxl.Workbook(); w = wb.active; w.title = 'README'
for i, l in enumerate(['Depth profiles of the net Poynting flux S_z (normalised to the flux returned to the device) with the manuscript optical constants.',
    'Stack: glass (n = 1.5) / ITO 50 or 150 nm (Koenig et al. 2014; 550 nm: 1.864 + 0.0032i) / organics 420 nm (n = 1.8) / Al (JO) or Ag (McPeak) 100 nm / air.',
    'Weighting: Ir(ppy)2acac spectrum, photon number, 400-700 nm (5 nm) x Lambertian cos.sin over 0-89 deg; s/p average.',
    "z = 0 at the glass/ITO interface. The flat value on the glass side is A' (all light not reflected); each step down is the loss in that layer (sheet summary).",
    'Source: figures/fig2cd_mock/export_profiles_konig.py (transfer-matrix code of fig2cd_mock.py with the library constants).'], 1): w.cell(i, 1, l)
w.column_dimensions['A'].width = 150
sm = wb.create_sheet('summary')
for j, h in enumerate(['case', "A' total (%)", 'ITO (%)', 'organics (%)', 'reflector (%)', 'leakage (%)'], 1): sm.cell(1, j, h).font = B
fig, axs = plt.subplots(1, 4, figsize=(16, 3.6), dpi=150)
for i, (name, L) in enumerate(cases):
    zabs, acc = weighted_profile(L); names = [l[0] for l in L]
    s = wb.create_sheet(name)
    for j, h in enumerate(['z (nm)', 'S_z / returned flux', 'layer'], 1): s.cell(1, j, h).font = B
    r = 2
    for j in range(len(zabs)):
        for zz, ss in zip(zabs[j], acc[j]): s.cell(r, 1, round(float(zz), 2)); s.cell(r, 2, float(ss)); s.cell(r, 3, names[j]); r += 1
    v = [name, 100*acc[0][0], 100*(acc[1][0]-acc[1][-1]), 100*(acc[2][0]-acc[2][-1]), 100*(acc[3][0]-acc[3][-1]), 100*acc[-1][0]]
    for j, x in enumerate(v, 1): sm.cell(i+2, j, x if j == 1 else round(float(x), 3))
    print(name, [round(float(x), 2) for x in v[1:]])
    ax = axs[i]; zz = np.concatenate(zabs); ss = np.concatenate(acc)
    for j in range(len(zabs)): ax.axvspan(zabs[j][0], zabs[j][-1], color=BAND.get(names[j], '#eee'), lw=0)
    ax.plot(zz, ss, color='#2a78d6' if 'Al' in name else '#eb6834', lw=2); ax.set_ylim(0, 0.2); ax.set_xlim(zz[0], zz[-1])
    ax.set_title(f"{name.replace('_', ', ')}: A' = {v[1]:.1f}% (ITO {v[2]:.1f}, metal {v[4]:.1f})", fontsize=8.5, loc='left')
    ax.set_xlabel('Depth from glass/ITO interface (nm)')
axs[0].set_ylabel('Net flux S_z / returned flux')
fig.tight_layout(); fig.savefig('fig2cd_profiles_konig.png')
wb.save('fig2cd_profiles_konig_rawdata.xlsx')
