# -*- coding: utf-8 -*-
"""Fig.1 mock-up: panel (a) full width on top, (b) and (c) side by side below.

This is the layout I am proposing, drawn so it can be judged rather than
described.  Everything except the two 3D renders is a native PowerPoint shape,
so any part of it can be nudged in PowerPoint itself.

Three things differ from the current draft on purpose:

* The cross-section's fills are *derived from the render's own band colours*
  (sampled at the label anchors), lightened by a fixed fraction so rays and
  stars still read on top.  Same layer, same colour, on both halves of the row.
  With the colours keyed that way the legend no longer needs swatches.
* One text size for every label and axis title, one for tick numbers, one for
  the panel letters.  Nothing else.
* The metal electrode is named once per row -- in the cross-section, where
  there is room for it -- so the 3D carries the four layers above it only, and
  the render is big enough to sit on the cross-section's baseline.

    python3 figures/make_full_figure_pptx.py
"""
import math, os, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import make_roundtrip_figure as F          # geometry + radiometry, single source

from pptx import Presentation
from pptx.util import Emu, Pt, Inches
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE as DASH
from pptx.oxml.ns import qn, nsdecls
from pptx.oxml import parse_xml

# ------------------------------------------------------------------ canvas
W, H = 1800.0, 1500.0                  # figure units; 1800 across = 180 mm printed
SLIDE_W = Inches(10.0)
SLIDE_H = Inches(10.0 * H / W)
SCALE = SLIDE_W / W                    # EMU per figure unit

def X(u): return int(round(u * SCALE))
Y = X
def D(u): return int(round(u * SCALE))
def PTS(u): return Pt(u * SCALE / 12700.0)

# text: three sizes in the whole figure and no more
T_LAB, T_TICK, T_LETTER = 25.0, 22.0, 31.0
T_3D = 20.0                            # inside the renders: bounded by the band width

# ------------------------------------------------------------------ palette
def C(h): return RGBColor.from_string(h.lstrip("#").upper())
def _rgb(h):
    h = h.lstrip("#"); return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))
def mix(h, k):                       # k = fraction of white blended in
    return "%02X%02X%02X" % tuple(int(round(c + k * (255 - c))) for c in _rgb(h))
def shade(h, k):
    return "%02X%02X%02X" % tuple(int(round(c * k)) for c in _rgb(h))

# measured off the renders at each label's anchor -- the 3D's own band colours
BAND = {"film": "D3D6D8", "glass": "B5BEC3", "tco": "A2CCD5", "org": "DBC599",
        "met_c": "8E97A4", "met_d": "A5A198"}
LENS_F  = "ECEFF1"                   # the film is lighter than the glass in the
GLASS_F = "D9E0E6"                   # render, and stays lighter here
TCO_F   = mix(BAND["tco"], 0.30)
ORG_F   = mix(BAND["org"], 0.30)
MET_F   = tuple(mix(BAND[k], 0.18) for k in ("met_c", "met_d"))   # per device,
                                     # as in the render, lifted just far enough
                                     # that dark ink reads on both
LINE_K  = 0.62                                # every outline is its fill, darkened

RAY, LOSS, EMIT, EMIT_EDGE = F.RAY, F.LOSS, F.EMIT, F.EMIT_EDGE
RAY_GLOW = "1FC79B"                  # the ray's hue, lifted: what a halo of it looks like
BURST = F.BURST
INK, MUTED, AXIS = "1A1A1A", "55585C", "2B2B2B"
BLU, GRN, RED = "3B6FD4", "3F9E4D", "DE5B5B"          # A' = 0.02 / 0.1 / 0.15
BAR = tuple(mix(c, 0.52) for c in (BLU, GRN, RED))
GREY = "8A8A8A"

# ------------------------------------------------------------------ layout
# One more round trip across a wider, shallower cross-section: the extra pass is
# what the panel is about, and spending the width on it instead of on scale is
# what lets (a) shrink.  Lens centres are PITCH*k + LENS_R and the zig-zag period
# is twice the glass-to-metal drop, so the hit lists stay on lens centres.
PW_F = 1200.0                                # 20 lenses, 5 lens hits, 4 bounces
HIT_TOP = [150.0, 390.0, 630.0, 870.0, 1110.0]
HIT_METAL = [270.0, 510.0, 750.0, 990.0]
NP = len(HIT_TOP)
X_EMIT = HIT_TOP[0] - (F.Y_EMIT - F.Y_TOP) * F.TAN

SX = 1.025                                   # cross-section scale, F units -> figure
CROSS_X = 530.0                              # 1230 wide -> right edge 1760
RW = 430.0                                   # 3D render, square
RX = 40.0
YB = (430.0, 800.0)                          # slab bottom of each row.  The pitch is
                                             # less than RW on purpose: a render's top
                                             # strip is glow below 8% alpha, so two may
                                             # overlap there and a row need not be as
                                             # tall as a square render
LEG_Y = 865.0
BX, BY, BW_, BH_ = 180.0, 988.0, 595.0, 380.0         # panel b plot box
CX_, CY_, CW_, CH_ = 1170.0, 988.0, 585.0, 380.0      # panel c plot box

P_ESC = 0.40                                 # escape probability per pass
A_SET = (0.02, 0.10, 0.15)                   # round-trip absorption, panel b
A_CONV, A_LOW = 0.15, 0.02                   # the two devices in panel a / c

prs = Presentation()
prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
slide = prs.slides.add_slide(prs.slide_layouts[6])
SH = slide.shapes

# ------------------------------------------------------------------ helpers
def _alpha(parent, a):
    if a is None or a >= 0.999 or parent is None:
        return
    sf = parent.find(qn("a:solidFill"))
    if sf is None:
        return
    clr = sf.find(qn("a:srgbClr"))
    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(a * 100000))}))

def style(s, fill=None, line=None, lw=None, op=None, name=None):
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = C(fill); _alpha(s._element.spPr, op)
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = C(line)
        s.line.width = Emu(max(3175, D(lw if lw is not None else 1.0)))
        _alpha(s._element.spPr.find(qn("a:ln")), op)
    s.shadow.inherit = False
    st = s._element.find(qn("p:style"))
    if st is not None:                 # the preset style's effectRef survives an
        s._element.remove(st)          # empty <a:effectLst/> in LibreOffice
    if name:
        s.name = name
    return s

def rect(x, y, w, h, fill, line=None, lw=1.2, op=None, name="rect"):
    return style(SH.add_shape(MSO_SHAPE.RECTANGLE, X(x), Y(y), D(w), D(h)),
                 fill, line, lw, op, name)

def oval(x, y, w, h, fill, line=None, lw=1.2, name="oval"):
    return style(SH.add_shape(MSO_SHAPE.OVAL, X(x), Y(y), D(w), D(h)),
                 fill, line, lw, None, name)

def seg(p0, p1, color, lw, head=False, op=None, dash=None, name="line"):
    cn = SH.add_connector(MSO_CONNECTOR.STRAIGHT, X(p0[0]), Y(p0[1]), X(p1[0]), Y(p1[1]))
    cn.line.color.rgb = C(color)
    cn.line.width = Emu(max(3175, D(lw)))
    ln = cn.line._get_or_add_ln()
    if head:
        ln.append(ln.makeelement(qn("a:tailEnd"),
                                 {"type": "triangle", "w": "lg", "len": "med"}))
    if dash:
        cn.line.dash_style = dash
    _alpha(ln, op)
    cn.shadow.inherit = False
    cn.name = name
    return cn

def poly(pts, color, lw, close=False, fill=None, dash=None, op=None, name="poly"):
    q = [(X(x), Y(y)) for x, y in pts]
    b = SH.build_freeform(q[0][0], q[0][1], scale=1.0)
    b.add_line_segments(q[1:], close=close)
    s = b.convert_to_shape()
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = C(fill)
    s.line.color.rgb = C(color)
    s.line.width = Emu(max(3175, D(lw)))
    if dash:
        s.line.dash_style = dash
    _alpha(s.line._get_or_add_ln(), op)
    s.shadow.inherit = False
    s.name = name
    return s

def wave(x0, y, dx, amp, periods, color, lw, op, head=False, name="wave"):
    n = 48
    pts = [(x0 + dx * i / n, y + amp * math.sin(2 * math.pi * periods * i / n))
           for i in range(n + 1)]
    s = poly(pts, color, lw, op=op, name=name)
    if head:
        ln = s.line._get_or_add_ln()
        ln.append(ln.makeelement(qn("a:tailEnd"),
                                 {"type": "triangle", "w": "med", "len": "med"}))
    return s

def star(cx, cy, r, kind, fill, line, op=None, lw=1.1, name="star"):
    s = SH.add_shape(kind, X(cx - r), Y(cy - r), D(2 * r), D(2 * r))
    try:
        s.adjustments[0] = 0.40
    except (IndexError, ValueError):
        pass
    return style(s, fill, line, lw, op, name)

def grad(s, stops, ang=90.0, centre=None):
    """Replace a shape's fill with a gradient.  `stops` = [(pos 0..1, hex, alpha)].

    Linear by default (DrawingML angle: 90 = top to bottom).  With `centre` =
    (fx, fy), fractions of the bounding box, it is a radial gradient whose first
    stop sits at that point -- used for glows that fall off from a source.
    """
    sp = s._element.spPr
    for t in ("a:solidFill", "a:noFill", "a:gradFill"):
        e = sp.find(qn(t))
        if e is not None:
            sp.remove(e)
    gs = "".join('<a:gs pos="%d"><a:srgbClr val="%s"><a:alpha val="%d"/></a:srgbClr></a:gs>'
                 % (int(p_ * 100000), c.lstrip("#").upper(), int(a * 100000))
                 for p_, c, a in stops)          # "#e8901f" is not a valid val: black
    if centre is None:
        tail = '<a:lin ang="%d" scaled="0"/>' % int(ang * 60000)
    else:
        l, t = int(centre[0] * 100000), int(centre[1] * 100000)
        tail = ('<a:path path="circle"><a:fillToRect l="%d" t="%d" r="%d" b="%d"/>'
                '</a:path>' % (l, t, 100000 - l, 100000 - t))
    el = parse_xml('<a:gradFill %s rotWithShape="1"><a:gsLst>%s</a:gsLst>%s</a:gradFill>'
                   % (nsdecls("a"), gs, tail))
    ln = sp.find(qn("a:ln"))
    if ln is not None:
        ln.addprevious(el)
    else:
        sp.append(el)
    return s

def glow(cx, cy, r, color, a0, name="glow", squash=1.0):
    """A soft radial glow: `color` at alpha a0 in the centre, nothing at radius r."""
    s = SH.add_shape(MSO_SHAPE.OVAL, X(cx - r), Y(cy - r * squash), D(2 * r),
                     D(2 * r * squash))
    style(s, "FFFFFF", None, name=name)
    return grad(s, [(0.0, color, a0), (0.45, color, a0 * 0.45), (1.0, color, 0.0)],
                centre=(0.5, 0.5))

def fan(ax, ay, r, half, color, a0, name="fan"):
    """Light leaving a lens: a sector opening upward from (ax, ay), fading out."""
    n = 24
    pts = [(ax, ay)] + [(ax + r * math.sin(math.radians(-half + 2 * half * i / n)),
                         ay - r * math.cos(math.radians(-half + 2 * half * i / n)))
                        for i in range(n + 1)]
    q = [(X(x), Y(y)) for x, y in pts]
    b = SH.build_freeform(q[0][0], q[0][1], scale=1.0)
    b.add_line_segments(q[1:], close=True)
    sh_ = b.convert_to_shape()
    style(sh_, "FFFFFF", None, name=name)
    # radial from the apex, which is the bottom-centre of the sector's box
    return grad(sh_, [(0.0, color, a0), (0.35, color, a0 * 0.55), (1.0, color, 0.0)],
                centre=(0.5, 1.0))

def text(cx, cy, runs, size, color, align=PP_ALIGN.CENTER, bold=False,
         w=600.0, rot=0.0, name="label", italic=False):
    """`runs` is a string, or [(text, baseline, italic), ...] for subscripts."""
    h = size * 1.7
    tb = SH.add_textbox(X(cx - w / 2.0), Y(cy - h / 2.0), D(w), D(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    if isinstance(runs, str):
        runs = [(runs, 0, italic)]
    for s_, base, ital in runs:
        r = p.add_run()
        r.text = s_
        r.font.size = PTS(size)
        r.font.bold = bold
        r.font.italic = ital
        r.font.name = "Arial"
        r.font.color.rgb = C(color)
        if base:
            r.font._rPr.set("baseline", "-25000" if base < 0 else "30000")
    if rot:
        tb.rotation = rot
    tb.name = name
    return tb

# ================================================================== panel a
def row(i, eta, r_met, t_tco, tag):
    yb, met = YB[i], MET_F[i]
    def fx(u): return CROSS_X + u * SX
    def fy(v): return yb + (v - F.Y_BOT) * SX
    def fd(d): return d * SX

    # --- micro-lens domes, then the slab over their lower halves
    # Each material gets a light falling from above: a white cap on every lens,
    # a brighter top to the glass, a specular strip along the metal's mirror
    # face.  Same base colours as before, so the match with the render holds.
    for k in range(int(PW_F / F.PITCH)):
        cx = F.PITCH * k + F.LENS_R
        grad(oval(fx(cx - F.LENS_R), fy(F.Y_TOP - F.LENS_R), fd(2 * F.LENS_R),
                  fd(2 * F.LENS_R), LENS_F, shade(LENS_F, LINE_K), 1.1,
                  "%s lens %02d" % (tag, k + 1)),
             [(0.0, "FFFFFF", 1.0), (0.30, "FFFFFF", 1.0), (0.50, LENS_F, 1.0),
              (1.0, LENS_F, 1.0)])
    grad(rect(fx(0), fy(F.Y_TOP), fd(PW_F), fd(F.Y_TCO - F.Y_TOP), GLASS_F,
              shade(GLASS_F, LINE_K), 1.3, name="%s glass" % tag),
         [(0.0, mix(GLASS_F, 0.55), 1.0), (1.0, GLASS_F, 1.0)])
    for y1, y2, fl, nm, st in (
            (F.Y_TCO, F.Y_ORG, TCO_F, "TCO",
             [(0.0, mix(TCO_F, 0.35), 1.0), (1.0, TCO_F, 1.0)]),
            (F.Y_ORG, F.Y_MET, ORG_F, "organic",
             [(0.0, mix(ORG_F, 0.30), 1.0), (1.0, ORG_F, 1.0)]),
            (F.Y_MET, F.Y_BOT, met, "metal",
             [(0.0, mix(met, 0.55), 1.0), (0.14, mix(met, 0.20), 1.0),
              (0.40, met, 1.0), (1.0, shade(met, 0.84), 1.0)])):
        grad(rect(fx(0), fy(y1), fd(PW_F), fd(y2 - y1), fl, shade(fl, LINE_K), 1.1,
                  name="%s %s" % (tag, nm)), st)

    # --- radiometry (same formulas as the SVG)
    I, sg, esc, loss = 1.0, [], [], []
    for k in range(NP):
        sg.append(I); esc.append(I * eta); I *= (1 - eta)
        if k < NP - 1:
            I *= t_tco; loss.append(I * (1 - r_met)); I *= r_met; I *= t_tco
    wid = [F.W_RAY * v ** 0.75 * SX for v in sg]
    kept = (1 - eta) ** 0.75

    tl = 1.0 - t_tco
    tw, ta = 10.0 + 180.0 * tl, 1.0 + 22.0 * tl
    for n, xm in enumerate(HIT_METAL):
        for j, sgn in enumerate((-1, 1)):
            wave(fx(xm + sgn * 24 - tw / 2), fy(F.Y_TCO + 8), fd(tw), fd(ta), 2,
                 LOSS, (0.9 + 9.0 * tl) * SX, min(1.0, 0.25 + 6.5 * tl),
                 name="%s TCO loss %d%s" % (tag, n + 1, "LR"[j]))

    # --- light leaving each lens: a fan whose strength is the power escaping
    #     there, so the eye reads the decay before it reads the arrow widths
    for k, xl in enumerate(HIT_TOP):
        fan(fx(xl), fy(F.Y_TOP), fd(108.0), 58.0, RAY_GLOW, 0.50 * sg[k] ** 0.8,
            "%s escape fan %d" % (tag, k + 1))
        glow(fx(xl), fy(F.Y_TOP - F.LENS_R * 0.45), fd(F.LENS_R * 1.15), RAY_GLOW,
             0.55 * sg[k] ** 0.8, "%s lit lens %d" % (tag, k + 1), squash=0.75)

    # --- a soft halo under every ray segment, same decay as the ray itself
    halo = [((X_EMIT + 7.1, F.Y_EMIT - 7.1), (HIT_TOP[0], F.Y_TOP), sg[0])]
    for k in range(NP - 1):
        halo.append(((HIT_TOP[k] + 9.2, F.Y_TOP + 9.2),
                     (HIT_METAL[k] - 6.4, F.Y_MET - 6.4), sg[k] * (1 - eta)))
        halo.append(((HIT_METAL[k] + 6.4, F.Y_MET - 6.4),
                     (HIT_TOP[k + 1], F.Y_TOP), sg[k + 1]))
    for n, (p0, p1, v) in enumerate(halo):
        seg((fx(p0[0]), fy(p0[1])), (fx(p1[0]), fy(p1[1])), RAY_GLOW,
            F.W_RAY * SX * 3.2 * v ** 0.6, op=0.20 * v ** 0.5,
            name="%s ray halo %d" % (tag, n + 1))

    glow(fx(X_EMIT), fy(F.Y_EMIT), fd(30.0), EMIT, 0.75, "%s emitter glow" % tag,
         squash=0.8)

    seg((fx(X_EMIT + 7.1), fy(F.Y_EMIT - 7.1)), (fx(HIT_TOP[0]), fy(F.Y_TOP)),
        RAY, wid[0], True, name="%s ray emit" % tag)
    for k in range(NP - 1):
        seg((fx(HIT_TOP[k] + 9.2), fy(F.Y_TOP + 9.2)),
            (fx(HIT_METAL[k] - 6.4), fy(F.Y_MET - 6.4)), RAY, wid[k] * kept, True,
            name="%s ray down %d" % (tag, k + 1))
        seg((fx(HIT_METAL[k] + 6.4), fy(F.Y_MET - 6.4)),
            (fx(HIT_TOP[k + 1]), fy(F.Y_TOP)), RAY, wid[k + 1], True,
            name="%s ray up %d" % (tag, k + 1))
    seg((fx(HIT_TOP[-1] + 9), fy(F.Y_TOP + 9)),
        (fx(PW_F), fy(F.Y_TOP + PW_F - HIT_TOP[-1])), RAY, wid[-1] * kept,
        False, 0.45, name="%s ray continues" % tag)

    for k, xl in enumerate(HIT_TOP):
        fw = max(1.0, 0.55 * wid[k])
        for j, (phi, ln_) in enumerate(((-30, 68), (-5, 79), (22, 70))):
            a, b = math.radians(phi), math.radians(phi * 1.45)
            o = (xl + F.LENS_R * math.sin(a), F.Y_TOP - F.LENS_R * math.cos(a))
            seg((fx(o[0]), fy(o[1])),
                (fx(o[0] + ln_ * math.sin(b)), fy(o[1] - ln_ * math.cos(b))),
                RAY, fw, True, name="%s escaping %d-%d" % (tag, k + 1, j + 1))

    for k, xm in enumerate(HIT_METAL):
        v = loss[k]
        glow(fx(xm), fy(F.Y_MET), fd(14.0 + 120.0 * v ** 0.7), LOSS,
             min(0.70, 4.0 * v), "%s heat %d" % (tag, k + 1), squash=0.55)
        # strictly proportional, no floor: a 2% loss should barely register
    for k, xm in enumerate(HIT_METAL):
        v = loss[k]
        r = 3.0 + 30.0 * v ** 0.6
        op = 0.5 + 0.5 * min(1.0, v / 0.18)
        for j, sgn in enumerate((-1, 1)):
            wave(fx(xm + sgn * (r - 1)), fy(F.Y_MET - 6), fd(sgn * (10.0 + 200.0 * v)),
                 fd(2.0 + 14 * v), 2.5, LOSS, (1.2 + 6.0 * v ** 0.6) * SX, op, True,
                 "%s ohmic loss %d%s" % (tag, k + 1, "LR"[j]))
        star(fx(xm), fy(F.Y_MET), fd(r), MSO_SHAPE.STAR_10_POINT, BURST, LOSS, op,
             1.0, "%s absorption %d" % (tag, k + 1))
    star(fx(X_EMIT), fy(F.Y_EMIT), fd(8.5), MSO_SHAPE.STAR_8_POINT, EMIT,
         EMIT_EDGE, None, 1.0, "%s emitter" % tag)
    oval(fx(X_EMIT) - D(2.4 * SX) / SCALE, fy(F.Y_EMIT) - D(2.4 * SX) / SCALE,
         fd(4.8), fd(4.8), "FFF8E6", None, name="%s emitter core" % tag)

    # --- the 3D render, bottom edge on the cross-section's baseline
    crop = F.render_crop(tag)
    ry = yb - RW
    if crop:
        pic = SH.add_picture(crop, X(RX), Y(ry + 0.015 * RW), D(RW), D(RW))
        pic.name = "%s render" % tag
        (bx, by, side), src_w = F._CROP[tag]
        k = RW / side
        for lab in F.render_labels_raw(tag):
            lx = RX + (lab["x"] * src_w - bx) * k
            ly = ry + 0.015 * RW + (lab["y"] * src_w - by) * k
            text(lx + 330.0, ly, lab["text"], T_3D,
                 "F2F5F8" if lab["light"] else INK, PP_ALIGN.LEFT,
                 w=660.0, name="%s: %s" % (tag, lab["text"]))
    return sum(esc)

# make_roundtrip_figure maps labels into its own slot; we need the raw fractions
def _raw(tag):
    import json
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                     "emitting_area_render_%s_labels.json" % tag)
    return json.load(open(p, encoding="utf-8")) if os.path.exists(p) else []
F.render_labels_raw = _raw

e1 = row(0, .30, .72, .90, "conventional")
e2 = row(1, .30, .98, .995, "designrule")

# ------------------------------------------------------------------ legend
style(SH.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, X(RX), Y(LEG_Y - 36), D(1760 - RX),
                   D(72)), "FFFFFF", "C9CED3", 1.0, None, "legend frame")
seg((RX + 55, LEG_Y), (RX + 130, LEG_Y), RAY, 7.0, True, name="legend ray")
text(RX + 148 + 150, LEG_Y, "Light ray", T_LAB, MUTED, PP_ALIGN.LEFT, w=300.0,
     name="legend ray text")
star(RX + 440, LEG_Y, 13.0, MSO_SHAPE.STAR_8_POINT, EMIT, EMIT_EDGE, None, 1.0,
     "legend emitter")
text(RX + 466 + 180, LEG_Y, "Exciton emission", T_LAB, MUTED, PP_ALIGN.LEFT,
     w=360.0, name="legend emission text")
wave(RX + 830, LEG_Y, 70, 4.4, 2.5, LOSS, 2.8, 0.95, True, "legend loss")
text(RX + 915 + 420, LEG_Y, "Absorption loss  (Joule dissipation at the metal or absorption in the TCO)",
     T_LAB, MUTED, PP_ALIGN.LEFT, w=840.0, name="legend loss text")

# ================================================================== axes kit
def frame(x, y, w, h, name):
    style(SH.add_shape(MSO_SHAPE.RECTANGLE, X(x), Y(y), D(w), D(h)),
          None, AXIS, 1.6, None, name)

def ticks_x(x, y, w, lo, hi, vals, fmt="%g", pad=16.0, name="x"):
    for v in vals:
        t = x + w * (v - lo) / (hi - lo)
        seg((t, y), (t, y - 11), AXIS, 1.4, name="%s tick" % name)
        text(t, y + pad + T_TICK * 0.5, fmt % v, T_TICK, INK, w=160.0,
             name="%s tick label" % name)

def ticks_y(x, y, h, lo, hi, vals, fmt="%g", side=-1, name="y"):
    for v in vals:
        t = y + h * (1 - (v - lo) / (hi - lo))
        seg((x, t), (x + side * -11, t), AXIS, 1.4, name="%s tick" % name)
        text(x + side * 24 + side * 60, t, fmt % v, T_TICK, INK,
             PP_ALIGN.RIGHT if side < 0 else PP_ALIGN.LEFT, w=120.0,
             name="%s tick label" % name)

# ================================================================== panel b
def panel_b():
    frame(BX, BY, BW_, BH_, "b frame")
    lo, hi = 0.0, 0.45
    ticks_x(BX, BY + BH_, BW_, -0.5, 10.5, range(11), "%d", name="b n")
    ticks_y(BX, BY, BH_, lo, hi, [0, .1, .2, .3, .4], "%.1f", -1, "b left")
    ticks_y(BX + BW_, BY, BH_, 0.0, 1.0, [0, .2, .4, .6, .8, 1.0], "%.1f", 1, "b right")
    text(BX + BW_ / 2, BY + BH_ + 86, "Number of round trips", T_LAB, INK,
         w=700.0, name="b x title")
    text(BX - 160, BY + BH_ / 2, "Escaped per pass", T_LAB, INK, w=500.0, rot=270.0,
         name="b y title")
    text(BX + BW_ + 165, BY + BH_ / 2, "Cumulative extraction", T_LAB, INK, w=520.0,
         rot=90.0, name="b y2 title")

    pitch = BW_ / 11.0
    bw = pitch * 0.78 / 3.0
    for j, a in enumerate(A_SET):
        r = (1 - P_ESC) * (1 - a)
        for n in range(11):
            v = P_ESC * r ** n
            x = BX + pitch * (n + 0.5) - pitch * 0.39 + bw * j
            hgt = BH_ * v / hi
            if hgt > 1.0:
                rect(x, BY + BH_ - hgt, bw, hgt, BAR[j], shade(BAR[j], 0.7), 0.8,
                     name="b bar A'=%.2f n=%d" % (a, n))
        pts = []
        for s in range(0, 121):
            n = 10.0 * s / 120.0
            cum = P_ESC * (1 - r ** (n + 1)) / (1 - r)
            pts.append((BX + pitch * (n + 0.5), BY + BH_ * (1 - cum)))
        poly(pts, (BLU, GRN, RED)[j], 2.6, name="b cumulative A'=%.2f" % a)

    # series key
    kx, ky = BX + BW_ * 0.60, BY + BH_ * 0.44
    text(kx + 60, ky - 34, [("A", 0, True), ("'", 0, False)], T_LAB, INK,
         PP_ALIGN.LEFT, w=120.0, name="b key title")
    for j, a in enumerate(A_SET):
        yy = ky + 34 * j
        rect(kx, yy - 12, 46, 24, BAR[j], shade(BAR[j], 0.7), 0.8,
             name="b key swatch %d" % j)
        seg((kx, yy + 14), (kx + 46, yy + 14), (BLU, GRN, RED)[j], 2.6,
            name="b key line %d" % j)
        text(kx + 66 + 60, yy, "%.2f" % a, T_LAB, INK, PP_ALIGN.LEFT, w=120.0,
             name="b key text %d" % j)
    text(BX + BW_ * 0.44, BY + BH_ * 0.46, [("p", 0, True), (" = 0.4", 0, False)],
         T_LAB, INK, w=200.0, name="b p label")
    # which axis each family belongs to
    seg((BX + BW_ * 0.30, BY + BH_ * 0.52), (BX + BW_ * 0.235, BY + BH_ * 0.655),
        INK, 1.8, True, name="b arrow bars")
    seg((BX + BW_ * 0.145, BY + BH_ * 0.45), (BX + BW_ * 0.215, BY + BH_ * 0.315),
        INK, 1.8, True, name="b arrow lines")

# ================================================================== panel c
def panel_c():
    frame(CX_, CY_, CW_, CH_, "c frame")
    lo, hi = 0.2, 1.0
    ticks_x(CX_, CY_ + CH_, CW_, 0.0, 1.0, [0, .2, .4, .6, .8, 1.0], "%.1f", name="c R")
    ticks_y(CX_, CY_, CH_, lo, hi, [.2, .4, .6, .8, 1.0], "%.1f", -1, "c eta")
    text(CX_ + CW_ / 2, CY_ + CH_ + 92,
         [("Round-trip reflectance, ", 0, False), ("R", 0, True),
          ("LED", -1, False), (" = 1 − ", 0, False), ("A", 0, True),
          ("'", 0, False)], T_LAB, INK, w=900.0, name="c x title")
    text(CX_ - 175, CY_ + CH_ / 2,
         [("Extraction efficiency (", 0, False), ("η", 0, True),
          ("sta", -1, False), (")", 0, False)], T_LAB, INK, w=600.0, rot=270.0,
         name="c y title")

    def eta(R, p): return p / (1.0 - (1.0 - p) * R)
    def pt(R, p): return (CX_ + CW_ * R, CY_ + CH_ * (1 - (eta(R, p) - lo) / (hi - lo)))

    for p, col, dsh in ((0.60, GREY, DASH.ROUND_DOT), (0.25, GREY, DASH.DASH),
                        (0.40, "E8252A", None)):
        poly([pt(s / 160.0, p) for s in range(161)], col, 2.6, dash=dsh,
             name="c curve p=%.2f" % p)
    text(CX_ + CW_ * 0.16, pt(0.16, 0.60)[1] - 34,
         [("p", 0, True), (" = 0.6", 0, False)], T_LAB, GREY, w=220.0, name="c p06")
    text(CX_ + CW_ * 0.32, pt(0.32, 0.25)[1] + 38,
         [("p", 0, True), (" = 0.25", 0, False)], T_LAB, GREY, w=240.0, name="c p025")
    text(CX_ + CW_ * 0.17, pt(0.17, 0.40)[1] + 34,
         [("p", 0, True), (" = 0.4", 0, False)], T_LAB, "E8252A", w=220.0, name="c p04")

    for a, col, lab in ((A_CONV, RED, "Conventional"), (A_LOW, BLU, "Low loss")):
        R = 1.0 - a
        x = CX_ + CW_ * R
        seg((x, CY_ + CH_), (x, pt(R, 0.40)[1]), col, 2.2, name="c marker %s" % lab)
        oval(x - 9, pt(R, 0.40)[1] - 9, 18, 18, col, "FFFFFF", 1.2,
             "c dot %s" % lab)
        # short enough to sit in the wedge below the p = 0.25 curve, so neither
        # note crosses a curve; (a) carries the full names
        text(x - 30, CY_ + CH_ * 0.72, lab, T_LAB, col, PP_ALIGN.LEFT, w=175.0,
             rot=270.0, name="c note %s" % lab)
        text(x - 62, CY_ + CH_ * 0.72,
             [("A", 0, True), ("' = %.2f" % a, 0, False)], T_LAB, col,
             PP_ALIGN.LEFT, w=175.0, rot=270.0, name="c value %s" % lab)

panel_b()
panel_c()

# ------------------------------------------------------------------ letters
for lx, ly, s in ((28, 30, "a"), (28, 948, "b"), (930, 948, "c")):
    text(lx + 30, ly, s, T_LETTER, INK, bold=True, w=90.0, name="panel %s" % s)

dest = os.path.join(os.path.dirname(os.path.abspath(__file__)), "figure1_layout.pptx")
prs.save(dest)
print("wrote %s\n  %d shapes, none grouped  |  extracted: %.0f%% vs %.0f%%"
      % (dest, len(SH), 100 * e1, 100 * e2))
