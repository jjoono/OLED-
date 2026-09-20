# -*- coding: utf-8 -*-
"""Fig. 2(i) on a fine grid: the chirped ZnS/LiF stack, angle every 0.25 deg and wavelength
every 1 nm, with Ag on the same grid for comparison."""
import numpy as np, openpyxl, os, sys
from openpyxl.styles import Font
sys.path.insert(0, os.path.join('..', '..', 'sim', 'design_rule4'))
import fig2d as F

BOLD = Font(bold=True)
NSUB = 1.5
LAM = np.arange(430.0, 701.0, 1.0)         # 271 rows
TH = np.radians(np.arange(0.0, 89.751, 0.25))   # 360 columns, grazing excluded (zero flux weight)
THC = np.degrees(np.arcsin(1/NSUB))

def rt(name):
    S = F.stacks()[name]; sel = np.isin(F.LAM, LAM)
    layers = [(m[sel], d) for m, d in S['layers']]
    return F.RT(layers, np.full(LAM.shape, NSUB, dtype=complex), S['exit'][sel], TH, LAM)

Rd, Td = rt('TCO + DBR')
Rg, Tg = rt('Ag')

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'README'
ws.column_dimensions['A'].width = 24; ws.column_dimensions['B'].width = 112
rows = [
 ('Fig. 2(i) on a fine grid — the chirped ZnS/LiF reflector', None),
 ('Angle every 0.25° from 0 to 89.75°, wavelength every 1 nm.  The coarse version in fig2ghij_rawdata.xlsx is 1° and 5 nm.  Exact grazing is left out: it carries zero flux weight and the transfer matrix is singular there.', None),
 (None, None), ('Stack', None),
 ('Structure', 'substrate (n = 1.50) / ITO 150 nm / 420 nm of non-absorbing organics (n = 1.8) / reflector, light arriving from the substrate'),
 ('reflector', 'an ITO 50 nm cathode on the organics — the transparent electrode a metal-free device needs — then 10 ZnS/LiF pairs whose thicknesses grow through the stack by 1.40: ZnS 55 → 77 nm, LiF 91 → 127 nm.'),
 ('Ag', 'McPeak measured n,k, 100 nm — on the same grid for comparison'),
 ('ITO', 'Koenig et al. 2014, full dispersion (1.8636 + 0.0032285i at 550 nm)'),
 (None, None), ('Is the grid fine enough?', None),
 ('resonance widths', 'at 550 nm the three features of the DBR are 1.30° wide at 40.8°, 2.43° at 70.3° and 2.65° at 82.5° (FWHM), so 0.25° puts five to ten points across each'),
 ('wavelength', 'at 60° the value moves by at most 0.40 %p per nm, 0.29 %p at the 95th percentile, so 1 nm is ample'),
 ('the one sharp edge', f'transmission switches off at the escape cone, {THC:.2f}°; beyond it the light is trapped by total internal reflection at the back surface and only absorption remains'),
 (None, None), ('Sheets', None),
 ('DBR_1mR', '100 × (1 − R), the round-trip loss — rows wavelength, columns angle'),
 ('DBR_T', '100 × T, the part transmitted out of the back.  Zero beyond the escape cone.'),
 ('DBR_absorbed', '100 × (1 − R − T), absorption alone.  This is the quantity to compare with Ag.'),
 ('Ag_absorbed', '100 × (1 − R − T) for the 100 nm Ag mirror on the same grid (its T is below 0.02 % everywhere)'),
 (None, None), ('For reference', None),
 ('weighted totals', 'flux- and spectrum-weighted over the green emission: ITO + DBR 4.08 % absorbed + 1.03 % leaked = 5.11 %; Ag 4.49 % + 0.01 % = 4.50 %.  With 15 pairs it reaches 4.67 %'),
 ('beyond the cone', 'nothing is transmitted there, so only absorption remains; with every layer made lossless it is identically zero, which is what pure total internal reflection requires'),
 (None, None), ('Provenance', None),
 ('Script', 'figures/fig2_full/export_i_fine.py, on sim/design_rule4/fig2d.py'),
 ('Repository', 'jjoono/OLED-, branch claude/oled-efficiency-paper-ag-955zsk'),
]
for i, (a, b) in enumerate(rows, start=1):
    if a is not None: ws.cell(row=i, column=1, value=a).font = BOLD
    if b is not None: ws.cell(row=i, column=2, value=b)
ws['A1'].font = Font(bold=True, size=14)

def sheet(title, note, M):
    s = wb.create_sheet(title)
    s['A1'] = note; s['A1'].font = BOLD
    s['A3'] = 'wavelength (nm) \\ angle (deg)'; s['A3'].font = BOLD
    deg = np.degrees(TH)
    for j in range(len(TH)):
        s.cell(row=3, column=2 + j, value=round(float(deg[j]), 2)).font = BOLD
    for i in range(len(LAM)):
        s.cell(row=4 + i, column=1, value=float(LAM[i])).font = BOLD
        row = M[i]
        for j in range(len(TH)):
            s.cell(row=4 + i, column=2 + j, value=round(float(row[j]), 4))
    s.column_dimensions['A'].width = 24
    s.freeze_panes = 'B4'

sheet('DBR_1mR', 'chirped DBR: 100 × (1 − R)', 100*(1 - Rd))
sheet('DBR_T', 'chirped DBR: 100 × T, the transmitted part', 100*Td)
sheet('DBR_absorbed', 'chirped DBR: 100 × (1 − R − T), absorption alone', 100*(1 - Rd - Td))
sheet('Ag_absorbed', 'Ag 100 nm: 100 × (1 − R − T), on the same grid', 100*(1 - Rg - Tg))
wb.save('fig2i_DBR_fine.xlsx')
print('saved fig2i_DBR_fine.xlsx', os.path.getsize('fig2i_DBR_fine.xlsx')//1024, 'kB')

hdr = 'rows 430:1:700 nm, cols 0:0.25:89.75 deg'
os.makedirs('ghij_csv', exist_ok=True)
for tag, M in (('DBR_1mR', 100*(1 - Rd)), ('DBR_T', 100*Td),
               ('DBR_absorbed', 100*(1 - Rd - Td)), ('Ag_absorbed', 100*(1 - Rg - Tg))):
    np.savetxt(f'ghij_csv/fine_{tag}.csv', M, delimiter=',', fmt='%.5f', header=hdr, comments='# ')
print('csv written')
