# -*- coding: utf-8 -*-
"""Raw data behind Fig. 2(g)-(j): the angle- and wavelength-resolved round-trip loss of the
three reflectors at a fixed 150 nm ITO, and the spectrum-averaged angle curves of (j)."""
import numpy as np, openpyxl, os, sys
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', 'sim', 'design_rule4'))
import fig2d as F

BOLD = Font(bold=True); YEL = PatternFill('solid', fgColor='FFFF00')
NSUB = 1.5
LAM_M = np.arange(430.0, 701.0, 5.0)          # map rows
TH_M = np.radians(np.arange(0.0, 90.0, 1.0))  # map columns
LAM_C = np.arange(430.0, 701.0, 1.0)          # curves: full spectral resolution
TH_C = np.radians(np.linspace(0.0, 89.9, 180))
SPEC = np.clip(np.interp(LAM_C, F.LAM, F.GREEN), 0, None); SPEC /= SPEC.sum()
PANELS = [('g', 'Al', 'Al'), ('h', 'Ag', 'Ag'), ('i', 'DBR, chirped', 'DBR (chirped, optimised)')]
CURVES = PANELS + [('j', 'DBR (10 pairs)', 'DBR (plain quarter-wave at 550 nm)')]

def rt(name, lam, th):
    S = F.stacks()[name]
    sel = np.isin(F.LAM, lam)
    layers = [(m[sel], d) for m, d in S['layers']]
    return F.RT(layers, np.full(lam.shape, NSUB, dtype=complex), S['exit'][sel], th, lam)

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'README'
ws.column_dimensions['A'].width = 26; ws.column_dimensions['B'].width = 108
rows = [
 ('Fig. 2(g)-(j) — round-trip loss of the three reflectors', None),
 ('Light arriving from the substrate on the OLED stack; 1 - R is what one round trip loses.', None),
 (None, None), ('Stack', None),
 ('Common part', 'substrate (n = 1.50) / ITO 150 nm / 420 nm of non-absorbing organics (n = 1.8) / reflector.  The transparent electrode is fixed, only the reflector changes.'),
 ('Al', 'Johnson-Christy-type n,k from the project library, 100 nm'),
 ('Ag', 'McPeak measured n,k, 100 nm'),
 ('DBR (chirped)', 'ZnS 56 -> 80 nm / LiF 88 -> 125 nm over 10 pairs, ZnS facing the organics.  Both thicknesses grow linearly through the stack by a factor 1.42, which widens the stopband instead of deepening it; the three numbers were chosen to minimise the flux- and spectrum-weighted loss over all substrate angles and the whole emission band.'),
 ('DBR (plain)', f'the same 10 pairs as a uniform quarter-wave stack at 550 nm: ZnS {F.D_ZNS:.1f} nm / LiF {F.D_LIF:.1f} nm.  Drawn dashed in (j); shown for comparison only.'),
 ('uniform stacks', 'for reference, the best uniform ZnS/LiF stacks give 5.46 / 5.72 / 6.00 % at 10 / 15 / 20 pairs — more pairs deepen the stopband but do not widen it, so they get worse.'),
 ('ITO', 'n = 1.8636 + 0.0032285i at 550 nm — Koenig et al., ACS Nano 8, 6182 (2014), full dispersion used'),
 (None, None), ('Quantities', None),
 ('1 - R', 'the round-trip loss: absorbed plus, for the dielectric mirror, transmitted out of the back'),
 ('T', 'the transmitted part alone.  Zero for the metals; for the DBR it is non-zero only inside the escape cone to air, theta < asin(1/n_sub) = 41.8 deg, since beyond it the light is trapped by total internal reflection at the back surface.'),
 ('weighting', 'angles are weighted cos(theta).sin(theta), the ergodic weight an angularly randomising outcoupling structure enforces; wavelengths by the green emission spectrum'),
 (None, None), ('Sheets', None),
 ('g_Al_map, h_Ag_map, i_DBR_map', '100(1 - R) as a matrix: rows are wavelength in 5 nm steps, columns are the angle in the substrate in 1 deg steps.  This is what the colour maps show.'),
 ('i_DBR_leak', 'the transmitted part 100T of the same DBR map, so the absorbed and leaked halves can be separated'),
 ('j_angle_curves', 'panel (j): 100(1 - R) averaged over the emission spectrum, against angle, for all four curves, plus the angular weight'),
 ('summary', 'the flux- and spectrum-weighted numbers quoted in the text'),
 (None, None), ('Provenance', None),
 ('Script', 'figures/fig2_full/export_ghij.py, on sim/design_rule4/fig2d.py (dispersive transfer-matrix, validated against the Octave dipole model to 1e-4)'),
 ('Repository', 'jjoono/OLED-, branch claude/oled-efficiency-paper-ag-955zsk'),
]
for i, (a, b) in enumerate(rows, start=1):
    if a is not None: ws.cell(row=i, column=1, value=a).font = BOLD
    if b is not None: ws.cell(row=i, column=2, value=b)
ws['A1'].font = Font(bold=True, size=14)

def write_matrix(title, note, M):
    s = wb.create_sheet(title)
    s['A1'] = note; s['A1'].font = BOLD
    s['A3'] = 'wavelength (nm) \\ angle (deg)'; s['A3'].font = BOLD
    for j, t in enumerate(np.degrees(TH_M)):
        s.cell(row=3, column=2 + j, value=float(t)).font = BOLD
    for i, l in enumerate(LAM_M):
        s.cell(row=4 + i, column=1, value=float(l)).font = BOLD
        for j in range(len(TH_M)):
            s.cell(row=4 + i, column=2 + j, value=float(M[i, j])).number_format = '0.00'
    s.column_dimensions['A'].width = 24
    s.freeze_panes = 'B4'

for _, name, label in PANELS:
    R, T = rt(name, LAM_M, TH_M)
    write_matrix(f'{_}_{"Al" if name=="Al" else "Ag" if name=="Ag" else "DBR"}_map',
                 f'{label}: 100 × (1 − R), rows wavelength / columns angle in the substrate', 100*(1 - R))
R, T = rt('DBR, chirped', LAM_M, TH_M)
write_matrix('i_DBR_leak', 'DBR (chirped): the transmitted part, 100 × T', 100*T)

s = wb.create_sheet('j_angle_curves')
s['A1'] = 'Panel (j): 100 × (1 − R) averaged over the green emission spectrum'; s['A1'].font = BOLD
s['A2'] = 'the weight column is cos(theta)sin(theta), normalised to its maximum; the escape cone to air ends at 41.81 deg'
hdr = ['angle in substrate (deg)'] + [lab for _, _, lab in CURVES] + ['cos.sin weight (norm.)']
for j, h in enumerate(hdr, start=1): s.cell(row=4, column=j, value=h).font = BOLD
cols = []
for _, name, _lab in CURVES:
    Rc, Tc = rt(name, LAM_C, TH_C)
    cols.append(100*(1 - (Rc*SPEC[:, None]).sum(0)))
w = np.cos(TH_C)*np.sin(TH_C); w /= w.max()
for i in range(len(TH_C)):
    s.cell(row=5 + i, column=1, value=float(np.degrees(TH_C[i]))).number_format = '0.0'
    for j, c in enumerate(cols):
        s.cell(row=5 + i, column=2 + j, value=float(c[i])).number_format = '0.000'
    s.cell(row=5 + i, column=2 + len(cols), value=float(w[i])).number_format = '0.0000'
s.column_dimensions['A'].width = 22
for j in range(2, 7): s.column_dimensions[get_column_letter(j)].width = 30
s.freeze_panes = 'A5'

s = wb.create_sheet('summary')
s['A1'] = 'Flux- and spectrum-weighted round-trip loss'; s['A1'].font = BOLD
s['A2'] = 'angles weighted cos.sin over 0-90 deg, wavelengths by the emission spectrum over 430-700 nm'
for j, h in enumerate(['reflector', 'absorbed (%)', 'transmitted (%)', 'total 1 − R (%)', 'substrate'], start=1):
    s.cell(row=4, column=j, value=h).font = BOLD
r = 5
for ns, spec, tag in ((1.5, F.GREEN, 'n_sub 1.50, green emitter'), (1.8, F.ORANGE, 'n_sub 1.80, orange emitter')):
    for _, name, lab in CURVES:
        W = F.weighted(name, n_sub=ns, spectrum=spec, lam=np.arange(430.0, 701.0))
        s.cell(row=r, column=1, value=lab)
        for j, v in enumerate((100*W['absorbed'], 100*W['transmitted'], 100*W['loss']), start=2):
            s.cell(row=r, column=j, value=float(v)).number_format = '0.00'
        s.cell(row=r, column=5, value=tag)
        r += 1
    r += 1
s.column_dimensions['A'].width = 34
for j in range(2, 6): s.column_dimensions[get_column_letter(j)].width = 19
wb.save('fig2ghij_rawdata.xlsx')
print('saved fig2ghij_rawdata.xlsx')

# full-resolution CSVs for the repository
os.makedirs('ghij_csv', exist_ok=True)
for _, name, lab in CURVES:
    R, T = rt(name, LAM_M, TH_M)
    tag = lab.split()[0].replace('(', '') + ('_plain' if 'plain' in lab else '')
    np.savetxt(f'ghij_csv/map_{tag}_1mR.csv', 100*(1 - R), delimiter=',',
               header='rows 430:5:700 nm, cols 0:1:89 deg, 100*(1-R)', comments='# ')
np.savetxt('ghij_csv/j_angle_curves.csv',
           np.column_stack([np.degrees(TH_C)] + cols + [w]), delimiter=',',
           header='angle_deg,' + ','.join(l.split(' (')[0].replace(' ', '_') for _, _, l in CURVES) + ',weight',
           comments='')
print('csv written')
