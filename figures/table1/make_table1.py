# -*- coding: utf-8 -*-
"""Table 1 of the manuscript as one pptx slide (journal-style three-line table), from sim/audit/table1_series.csv."""
import csv
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

rows = list(csv.DictReader(open('/home/user/OLED-/sim/audit/table1_series.csv')))
prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
s = prs.slides.add_slide(prs.slide_layouts[6])
F = 'Arial'

def text(x, y, w, h, runs, size, color='000000', anchor=MSO_ANCHOR.TOP):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0; tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    for t, bold in runs:
        r = p.add_run(); r.text = t; r.font.name = F; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = RGBColor.from_string(color)
    return tb

text(0.6, 0.35, 12.1, 0.45, [('표 1. ', True), ('실제 재료 스택으로 계산한 near-unity 설계의 EQE (550 nm, PLQY = 1)', False)], 16)
text(0.6, 0.85, 12.1, 0.8, [('스택: Ag 100 nm / B3PyMPM d_ETL / TCTA:B3PyMPM 25 nm (발광 위치 중앙) / TAPC 180 nm / ITO 50 nm / 기판 n = 1.8 + 반구형 microlens array (n_MLA = 1.8, 행렬 급수). '
                             '광학상수는 실측값(B3PyMPM n_o/n_e = 1.821/1.609, TCTA:B3PyMPM 1.833/1.671, TAPC 1.691/1.664); ITO는 König 등의 n = 1.864에 k = 0.0032(문헌값) 또는 0.002.', False)], 11, '404040')

hdr = ['d_ETL (nm)', 'k_ITO', 'Θ', 'η_sub', 'SPP 손실', 'A′', 'η_ext', 'EQE']
colw = [1.35, 1.2, 1.55, 1.5, 1.5, 1.5, 1.75, 1.75]
tbl = s.shapes.add_table(len(rows) + 1, 8, Inches(0.6), Inches(1.8), Inches(sum(colw)), Inches(0.32 * (len(rows) + 1))).table
for j, w in enumerate(colw): tbl.columns[j].width = Inches(w)
tblPr = tbl._tbl.tblPr; tblPr.set('firstRow', '0'); tblPr.set('bandRow', '0')
for st in tblPr.findall(qn('a:tableStyleId')): tblPr.remove(st)

def border(cell, side, pt=None, color='000000'):
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(qn('a:' + side)): tcPr.remove(old)
    ln = etree.SubElement(tcPr, qn('a:' + side))
    if pt is None: etree.SubElement(ln, qn('a:noFill'))
    else:
        ln.set('w', str(int(pt * 12700))); sf = etree.SubElement(ln, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', color)

def setcell(cell, t, bold=False, fill='FFFFFF', size=12, top=None, bottom=None, bcol='000000'):
    cell.fill.solid(); cell.fill.fore_color.rgb = RGBColor.from_string(fill)
    cell.margin_left = cell.margin_right = Inches(0.05); cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame; p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = t; r.font.name = F; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = RGBColor(0, 0, 0)
    border(cell, 'lnL'); border(cell, 'lnR'); border(cell, 'lnT', top, bcol); border(cell, 'lnB', bottom, bcol)

for j, h in enumerate(hdr): setcell(tbl.cell(0, j), h, bold=True, top=1.25, bottom=0.75)
for i, r in enumerate(rows):
    th = float(r['Theta']); ths = '0.67 (등방성)' if abs(th - 0.667) < 1e-3 else '%.2f' % th
    vals = [r['d_ETL_nm'], '%.4f' % float(r['k_ITO']), ths, '%.3f' % float(r['eta_sub']), '%.3f' % float(r['spp']),
            '%.3f' % float(r['Aprime']), '%.3f' % float(r['eta_ext_series']), '%.3f' % float(r['EQE_series'])]
    last = i == len(rows) - 1; grp = (i + 1) % 3 == 0 and not last
    fill = 'F2F2F2' if (3 <= i < 6 or i >= 9) else 'FFFFFF'
    for j, v in enumerate(vals):
        setcell(tbl.cell(i + 1, j), v, bold=(j == 7), fill=fill, bottom=(1.25 if last else (0.5 if grp else None)), bcol=('000000' if last else 'A0A0A0'))
for i in range(len(rows) + 1): tbl.rows[i].height = Inches(0.32)

ec = [float(r['eta_ext_closed']) for r in rows]; qc = [float(r['EQE_closed']) for r in rows]; qs = [float(r['EQE_series']) for r in rows]
text(0.6, 6.15, 12.1, 0.7, [('회색 행은 흡수가 더 작은 ITO(k = 0.002). Θ = 0.67은 등방성. SPP 손실은 evanescent 성분 전체(SPP 모드 포함). 식 (3)으로 계산하면 η_ext %.3f–%.3f, EQE %.3f–%.3f로 급수보다 %.1f–%.1f %%p 높다(식 (3)은 상한). 원자료: sim/audit/table1_series.csv.' % (min(ec), max(ec), min(qc), max(qc), 100 * min(q - s_ for q, s_ in zip(qc, qs)), 100 * max(q - s_ for q, s_ in zip(qc, qs))), False)], 10.5, '404040')
prs.save('table1.pptx'); print('table1.pptx written')
