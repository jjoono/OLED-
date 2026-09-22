# -*- coding: utf-8 -*-
"""One workbook with exactly the columns to redraw every panel whose numbers changed:
Fig. 2(c)-(f), Fig. 3(a), Fig. 5(c), the new Fig. 5(d), and Supplementary Table 1."""
import numpy as np, openpyxl, os
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
R = '/home/user/OLED-'
def load(fn): return np.genfromtxt(fn, delimiter=',', names=True, dtype=None, encoding='utf-8')
BOLD = Font(bold=True)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'README'
ws.column_dimensions['A'].width = 18; ws.column_dimensions['B'].width = 118
rows = [
 ('Panels to redraw', 'Fig. 2(c)-(f): closed form eq. (3), eta_ext = p / [p + (1 - p) A\'], with p the cos.sin-weighted B_T of the hemispherical MLA (n_MLA = n_sub) and A\' the cos.sin-weighted 1 - R_LED.  Fig. 3(a), 5(c), 5(d), SI Table 1: matrix series of eq. (1).  Decision of 2026-09-22: the series shows mode cut-off steps against n_sub (SI_Fig1_nsub), so the trend panels of Fig. 2 use eq. (3).'),
 (None, None),
 ('Fig. 2(c)', 'eta_ext vs ITO thickness, n_sub 1.5, eq. (3) with p = 0.427.  Ag 0.969 (30 nm) -> 0.933 (200 nm), Al 0.843 -> 0.822; Ag above Al by 11-13 %p.  Measured green device 0.916 at 150 nm ITO for reference.'),
 ('Fig. 2(d)', 'p and A\' vs n_sub.  p is now the structure property: 0.58 (1.3) -> 0.43 (1.5) -> 0.29 (1.8) -> 0.23 (2.0), 2.5x; was the device-weighted 0.665 -> 0.248.  A\' unchanged.'),
 ('Fig. 2(e)', 'eta_ext vs n_sub, eq. (3).  Gap Ag - Al: 7 %p (1.3), 12 %p (1.5), 22 %p (2.0).  Smooth; the series version with its cut-off dips is in SI_Fig1_nsub.'),
 ('Fig. 2(f)', 'eta_sub (unchanged) and EQE vs n_sub, eq. (3); extraction loss = eta_sub - EQE.  Al 61 % maximum at 1.8 then 58-60 %; Ag 89 % at 1.8, 87 % at 2.0.'),
 ('SI_Fig1_nsub', 'Supplementary Fig. 1 data: same stack, n_sub in 0.01 steps (BSDF interpolated between the 0.05 slices): eta_ext and EQE by eq. (3) and by the series, P_sub fraction beyond 70 deg.  Guided modes at 550 nm: n_eff 1.352 (TM), 1.523 (TE), 1.663 (TM), 1.734 (TE).'),
 ('Fig. 2(a),(b),(g)-(j)', 'unchanged (A\' split and the angle/wavelength maps do not depend on the extraction model).'),
 ('Fig. 3(a)', 'k_ITO and n_Ag sweeps at n_sub 1.8, series.  k 0 -> 0.01: eta_ext 0.95 -> 0.81, EQE 0.91 -> 0.75 (was 0.96 -> 0.88, 0.93 -> 0.82); k 0.08: EQE 0.45 (was 0.49).  n_Ag 0.2: EQE 0.55 (was 0.60).'),
 ('Fig. 3(b),(c)', 'unchanged.'),
 ('Fig. 5(c)', 'EQE vs Theta, series.  proposed 83.1 -> 85.0 % (was 88.4 -> 91.2), reference 31.3 -> 54.8 % (was 34 -> 60).'),
 ('Fig. 5(d)', 'new panel: EQE vs lens aspect ratio, n_sub = n_MLA 1.8, close-packed hexagonal, series.  Ag max 86.6 % at AR 0.55, >= 95 % of it for AR 0.25-1.5; Al max 56.2 % at 0.6, from 0.35.  Flat interface 25.5 / 21.6 %.  The AR = 1 row equals the LightTools slice-11 hemisphere.'),
 ('Fig. 5(a),(b),(e), Fig. 4', 'unchanged (author\'s series data).'),
 ('SI Table 1', 'real material stack, series: 87 % isotropic, 87.5-88 % at Theta 0.8-0.9, 89-90 % with k_ITO 0.002.  Eq. (3) values in brackets for reference.'),
 (None, None),
 ('Stack', 'air | reflector 100 nm | ETL 1.8 | EML 1.8 20 nm | HTL 1.8 | ITO (Koenig 2014) | n_sub; 550 nm, PLQY 1, Theta 2/3 unless swept; Ag McPeak, Al Johnson-Christy type.  Fig. 2(d)-(f) and 3(a): ITO 50 nm; Fig. 2(c): ITO swept; Fig. 5(c) reference: Al, n 1.5, ETL/HTL 80/230; proposed: Ag, n 1.8, 200/200.'),
 ('Source', 'sim/audit/recompute_series.py, sim/mla/run_aspect.py; repository jjoono/OLED-, branch claude/oled-efficiency-paper-ag-955zsk'),
]
for i, (a, b) in enumerate(rows, 1):
    ws.cell(i, 1, a).font = BOLD if a else Font(); ws.cell(i, 2, b)

def sheet(name, cols, rows_):
    s = wb.create_sheet(name)
    for j, h in enumerate(cols, 1):
        s.cell(1, j, h).font = BOLD; s.column_dimensions[get_column_letter(j)].width = max(13, len(h) + 2)
    for i, r in enumerate(rows_, 2):
        for j, v in enumerate(r, 1):
            c = s.cell(i, j, v)
            if isinstance(v, float): c.number_format = '0.0000'
    s.freeze_panes = 'A2'

c = load(R + '/sim/audit/fig2c_series.csv')
al, ag = c[c['reflector'] == 'Al'], c[c['reflector'] == 'Ag']
sheet('Fig2c', ['d_ITO_nm', 'Aprime_Al', 'eta_ext_Al', 'Aprime_Ag', 'eta_ext_Ag'],
      [[float(a['d_ITO_nm']), float(a['Aprime']), float(a['eta_ext_closed']), float(g['Aprime']), float(g['eta_ext_closed'])] for a, g in zip(al, ag)])
d = load(R + '/sim/audit/fig2def_series.csv')
al, ag = d[d['reflector'] == 'Al'], d[d['reflector'] == 'Ag']
sheet('Fig2d', ['n_sub', 'p_structure', 'Aprime_Al', 'Aprime_Ag'], [[float(a['n_sub']), float(a['p_cos_sin']), float(a['Aprime']), float(g['Aprime'])] for a, g in zip(al, ag)])
sheet('Fig2e', ['n_sub', 'eta_ext_Al', 'eta_ext_Ag'], [[float(a['n_sub']), float(a['eta_ext_closed']), float(g['eta_ext_closed'])] for a, g in zip(al, ag)])
sheet('Fig2f', ['n_sub', 'eta_sub_Al', 'EQE_Al', 'extraction_loss_Al', 'eta_sub_Ag', 'EQE_Ag', 'extraction_loss_Ag'],
      [[float(a['n_sub']), float(a['eta_sub']), float(a['EQE_closed']), float(a['eta_sub'] - a['EQE_closed']), float(g['eta_sub']), float(g['EQE_closed']), float(g['eta_sub'] - g['EQE_closed'])] for a, g in zip(al, ag)])
fn = load(R + '/sim/audit/fig2def_fine.csv'); fal, fag = fn[fn['reflector'] == 'Al'], fn[fn['reflector'] == 'Ag']
sheet('SI_Fig1_nsub', ['n_sub', 'eta_ext_eq3_Al', 'eta_ext_series_Al', 'EQE_eq3_Al', 'EQE_series_Al', 'Psub_beyond70deg_Al', 'eta_ext_eq3_Ag', 'eta_ext_series_Ag', 'EQE_eq3_Ag', 'EQE_series_Ag', 'Psub_beyond70deg_Ag'],
      [[float(a['n_sub']), float(a['eta_ext_closed']), float(a['eta_ext_series']), float(a['EQE_closed']), float(a['EQE_series']), float(a['graze70']), float(g['eta_ext_closed']), float(g['eta_ext_series']), float(g['EQE_closed']), float(g['EQE_series']), float(g['graze70'])] for a, g in zip(fal, fag)])
f = load(R + '/sim/audit/fig3a_series.csv')
for sw, name in (('ITO_k', 'Fig3a_kITO'), ('Ag_n', 'Fig3a_nAg')):
    m = f[f['sweep'] == sw]
    sheet(name, ['k_ITO' if sw == 'ITO_k' else 'n_Ag', 'eta_sub', 'eta_ext', 'EQE'], [[float(r['value']), float(r['eta_sub']), float(r['eta_ext_series']), float(r['EQE_series'])] for r in m])
t = load(R + '/sim/audit/fig5c_series.csv')
ref, pro = t[t['device'] == 'reference'], t[t['device'] == 'proposed']
sheet('Fig5c', ['Theta', 'EQE_reference', 'eta_sub_reference', 'EQE_proposed', 'eta_sub_proposed'],
      [[float(a['Theta']), float(a['EQE_series']), float(a['eta_sub']), float(b['EQE_series']), float(b['eta_sub'])] for a, b in zip(ref, pro)])
A = load(R + '/sim/mla/mla_aspect.csv'); A = A[A['aspect_ratio'] >= 0.09]
sheet('Fig5d', ['aspect_ratio', 'p_single_pass', 'EQE_Al', 'EQE_Ag', 'lens_shape'], [[float(r['aspect_ratio']), float(r['p_single_pass']), float(r['EQE_Al']), float(r['EQE_Ag']), str(r['lens_shape'])] for r in A])
F = load(R + '/sim/mla/mla_flat_reference.csv'); L = load(R + '/sim/mla/mla_aspect_lt_slice11.csv')
sheet('Fig5d_reference_lines', ['item', 'Al', 'Ag'],
      [['flat interface EQE', float(F[F['reflector'] == 'Al']['EQE_flat'][0]), float(F[F['reflector'] == 'Ag']['EQE_flat'][0])],
       ['LightTools slice 11 hemisphere EQE', float(L[L['reflector'] == 'Al']['EQE'][0]), float(L[L['reflector'] == 'Ag']['EQE'][0])]])
s1 = load(R + '/sim/audit/table1_series.csv')
sheet('SI_Table1', ['d_ETL_nm', 'k_ITO', 'Theta', 'eta_sub', 'SPP', 'Aprime', 'eta_ext_series', 'EQE_series', 'eta_ext_eq3', 'EQE_eq3'],
      [[float(r['d_ETL_nm']), float(r['k_ITO']), float(r['Theta']), float(r['eta_sub']), float(r['spp']), float(r['Aprime']), float(r['eta_ext_series']), float(r['EQE_series']), float(r['eta_ext_closed']), float(r['EQE_closed'])] for r in s1])
wb.save('figure_update_rawdata.xlsx'); print('saved')
