# -*- coding: utf-8 -*-
"""Raw data behind the Fig. 2(b) / 2(c) square panels, as a workbook.
eta_ext and EQE are live formulas driven by the editable p column, so the panels
can be redrawn with a different escape-probability curve without rerunning the model."""
import numpy as np, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BOLD = Font(bold=True)
YEL = PatternFill('solid', fgColor='FFFF00')
F4, F2 = '0.0000', '0.00'

def pflat(n):
    """Flat-interface escape probability for angularly randomised substrate light."""
    thc = np.arcsin(1/n)
    e = np.linspace(0, thc, 20001); th = 0.5*(e[1:]+e[:-1])
    st = n*np.sin(th); ct = np.sqrt(1-st**2); cs = np.cos(th)
    rs = (n*cs - ct)/(n*cs + ct); rp = (n*ct - cs)/(n*ct + cs)
    T = 1 - 0.5*(rs**2 + rp**2)
    return 2*np.trapezoid(T*np.sin(th)*np.cos(th), th)

C = np.genfromtxt('fig2b_curves.csv', delimiter=',', names=True)
AL = np.genfromtxt('nsub_al.csv', delimiter=',')
AG = np.genfromtxt('nsub_ag.csv', delimiter=',')
n = C['n_sub']; N = len(n)

wb = openpyxl.Workbook()

# ------------------------------------------------------------------ README
ws = wb.active; ws.title = 'README'
ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 104
rows = [
 ('Fig. 2(b) and 2(c) — substrate / outcoupling-structure index sweep', None),
 ('Raw data for the two square panels: (b) p and eta_ext, (c) eta_sub and EQE, each for an Al and an Ag reflector', None),
 (None, None),
 ('Panels', None),
 ('(b)', 'eta_ext against n_sub for both reflectors, with the single-pass escape probability p on the same axes'),
 ('(c)', 'eta_sub (dashed) and EQE = eta_sub x eta_ext (solid) against n_sub for both reflectors'),
 (None, None),
 ('Stack', None),
 ('Common part', 'reflector 100 nm / ETL 200 nm / EML 20 nm (isotropic dipole at the centre) / HTL 200 nm / ITO 50 nm / substrate'),
 ('Reflector', 'Al: Johnson-Christy-type n,k from the project library (0.958 + 6.687i at 550 nm).  Ag: McPeak measured n,k (0.044 + 3.819i)'),
 ('ITO', 'n = 1.9 + 0.02i, fixed at 50 nm'),
 ('Organics', 'EML, HTL and ETL isotropic n = 1.8, k = 0'),
 ('Substrate', 'index swept 1.30 to 2.00 in 0.05 steps, index-matched to the outcoupling structure'),
 ('Wavelength', '550 nm, single wavelength, PLQY = 1, u grid 3000 points'),
 (None, None),
 ('Quantities', None),
 ('eta_sub', 'substrate-delivered power from the five-channel budget: 1 minus waveguided, SPP and absorbed fractions'),
 ("A'", 'round-trip loss = 1 minus the reflectance of the OLED stack seen from the substrate, flux-weighted over substrate angles and averaged over p and s'),
 ('p', 'single-pass escape probability of the outcoupling structure, from the ray trace.  EDITABLE — column B of sheet fig2b_2c drives every eta_ext and EQE formula'),
 ('eta_ext', "p / [p + (1-p) A'].  FORMULA"),
 ('EQE', 'eta_sub x eta_ext.  FORMULA'),
 ('p (flat interface)', 'reference column only, not used by any formula: the escape probability of a flat substrate/air interface for angularly randomised light, mean Fresnel transmittance over the escape cone divided by n_sub^2 (Yablonovitch 1982, one-sided)'),
 (None, None),
 ('Sheets', None),
 ('fig2b_2c', 'the plotted curves, with eta_ext and EQE as formulas'),
 ('Al_modes', 'the full five-channel power budget for the Al reflector, as the model returns it'),
 ('Ag_modes', 'the same for the Ag reflector'),
 (None, None),
 ('Provenance', None),
 ('Script', 'sim/design_rule4/dr4f.m (Octave), MODE=nsub, run with UNUM=3000 and TOPMAT=al / ag'),
 ('Source files', 'nsub_al.csv and nsub_ag.csv (model output), fig2b_curves.csv (the assembled curves)'),
 ('p', "interpolated from the author's single-pass escape-probability curve: 0.665 / 0.470 / 0.380 / 0.355 / 0.335 / 0.305 / 0.275 / 0.248 at n_sub = 1.30 / 1.40 / 1.50 / 1.60 / 1.70 / 1.80 / 1.90 / 2.00"),
 ('Caution', "column 9 of nsub_al.csv and nsub_ag.csv is eta_ext at the script's default p = 0.4 and is NOT the plotted curve; the workbook recomputes eta_ext from the p column"),
 ('Repository', 'jjoono/OLED-, branch claude/oled-efficiency-paper-ag-955zsk, figures/fig2b_mock/'),
]
for i, (a, b) in enumerate(rows, start=1):
    if a is not None:
        ws.cell(row=i, column=1, value=a).font = BOLD
    if b is not None:
        c = ws.cell(row=i, column=2, value=b); c.alignment = Alignment(wrap_text=False)
ws['A1'].font = Font(bold=True, size=14)

# ------------------------------------------------------------- plotted curves
ws = wb.create_sheet('fig2b_2c')
ws['A1'] = 'The curves as plotted in fig2bc_squares.png'; ws['A1'].font = BOLD
ws['A2'] = "eta_ext = p/[p+(1-p)A'] and EQE = eta_sub x eta_ext are formulas; edit column B (yellow) to redraw with a different p"
hdr = ['n_sub', 'p', 'eta_sub (Al)', "A' (Al)", 'eta_ext (Al)', 'EQE (Al)',
       'eta_sub (Ag)', "A' (Ag)", 'eta_ext (Ag)', 'EQE (Ag)',
       'eta_ext gap (Ag - Al)', 'p, flat interface (reference)']
for j, h in enumerate(hdr, start=1):
    ws.cell(row=4, column=j, value=h).font = BOLD
for i in range(N):
    r = 5 + i
    ws.cell(row=r, column=1, value=float(n[i])).number_format = F2
    c = ws.cell(row=r, column=2, value=float(C['p'][i])); c.number_format = F4; c.fill = YEL; c.font = BOLD
    ws.cell(row=r, column=3, value=float(C['eta_sub_Al'][i])).number_format = F4
    ws.cell(row=r, column=4, value=float(C['Aprime_Al'][i])).number_format = F4
    ws.cell(row=r, column=5, value=f'=$B{r}/($B{r}+(1-$B{r})*D{r})').number_format = F4
    ws.cell(row=r, column=6, value=f'=C{r}*E{r}').number_format = F4
    ws.cell(row=r, column=7, value=float(C['eta_sub_Ag'][i])).number_format = F4
    ws.cell(row=r, column=8, value=float(C['Aprime_Ag'][i])).number_format = F4
    ws.cell(row=r, column=9, value=f'=$B{r}/($B{r}+(1-$B{r})*H{r})').number_format = F4
    ws.cell(row=r, column=10, value=f'=G{r}*I{r}').number_format = F4
    ws.cell(row=r, column=11, value=f'=I{r}-E{r}').number_format = F4
    ws.cell(row=r, column=12, value=float(pflat(n[i]))).number_format = F4
ws.column_dimensions['A'].width = 8
for j in range(2, 13):
    ws.column_dimensions[get_column_letter(j)].width = 21 if j in (11, 12) else 14
ws.freeze_panes = 'A5'

# ----------------------------------------------------------- the mode budgets
for tag, D, name in (('Al', AL, 'Al_modes'), ('Ag', AG, 'Ag_modes')):
    ws = wb.create_sheet(name)
    ws['A1'] = f'Five-channel power budget, {tag} reflector — as dr4f.m returns it'
    ws['A1'].font = BOLD
    ws['A2'] = 'Columns B to E sum to 1 at PLQY = 1 (column J is the check).  F and G split the absorbed fraction by layer.'
    hdr = ['n_sub', 'eta_sub', 'waveguided', 'SPP', 'absorbed (total)',
           '  in the reflector', '  in the TCO / organics', "A' (round trip)",
           'eta_ext at p = 0.4', 'sum check']
    for j, h in enumerate(hdr, start=1):
        ws.cell(row=4, column=j, value=h).font = BOLD
    for i in range(N):
        r = 5 + i
        vals = [D[i, 0], D[i, 1], D[i, 3], D[i, 4], D[i, 5], D[i, 6], D[i, 7], D[i, 2], D[i, 8]]
        for j, v in enumerate(vals, start=1):
            ws.cell(row=r, column=j, value=float(v)).number_format = F2 if j == 1 else F4
        ws.cell(row=r, column=10, value=f'=B{r}+C{r}+D{r}+E{r}').number_format = F4
    ws.column_dimensions['A'].width = 8
    for j in range(2, 11):
        ws.column_dimensions[get_column_letter(j)].width = 22 if j in (6, 7) else 17
    ws.freeze_panes = 'A5'

wb.save('fig2bc_rawdata.xlsx')
print('saved fig2bc_rawdata.xlsx')
