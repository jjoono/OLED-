"""Figures + Excel data for the thin-film collimator study (reads results.json, blur.py)."""
import json, numpy as np, matplotlib
matplotlib.use('Agg'); import matplotlib.pyplot as plt
from openpyxl import Workbook
import model as Mo, blur

R = json.load(open('results.json'))
rows = R['rows']; BETAS = R['beta']; FWHMS = R['fwhm']
plt.rcParams.update({'font.size': 7, 'axes.titlesize': 7.5})
C = {10.0: '#2E7D5B', 20.0: '#3B6EA5', 30.0: '#8E5BA6', 50.0: '#C6512F'}
AF = (1.5, 1.6)          # OLED lifetime acceleration factor range (LT ~ L^-n)
PANCAKE = 0.125          # unpolarized pancake ceiling (half-mirror x polarizer)


def get(f, a, b):
    return next(r for r in rows if r['fwhm'] == f and r['a_sc'] == a and r['beta'] == b)


def save(fig, name, sheets):
    fig.savefig(name + '.png', dpi=300, bbox_inches='tight'); fig.savefig(name + '.pdf', bbox_inches='tight')
    wb = Workbook(); wb.remove(wb.active)
    for title, header, data in sheets:
        ws = wb.create_sheet(title[:31]); ws.append(header)
        for d in data: ws.append([float(x) if isinstance(x, (np.floating, float, int, np.integer)) else x for x in d])
    wb.save(name + '.xlsx'); print('saved', name)


# ---- Fig 1: power into the acceptance cone vs beta ----
fig, ax = plt.subplots(1, 2, figsize=(7.2, 2.7), gridspec_kw=dict(wspace=0.32))
sh = []
for f in FWHMS:
    eb = [get(f, 0.0, b)['eta_beta'] for b in BETAS]
    ax[0].plot(BETAS, eb, 'o-', color=C[f], ms=3, label=f'scatterer+DBR, FWHM {f:.0f} nm')
    sh.append((f'DBR_FWHM{f:.0f}', ['beta', 'eta_beta', 'eta_total', 'mat', 'N', 'lam_b'],
               [[b, get(f, 0.0, b)['eta_beta'], get(f, 0.0, b)['eta_total'], get(f, 0.0, b)['mat'], get(f, 0.0, b)['N'], get(f, 0.0, b)['lam_b']] for b in BETAS]))
ff = 30.0
ax[0].plot(BETAS, [get(ff, 0.0, b)['ideal_eta_beta'] for b in BETAS], 'k--', lw=1, label='ideal angular filter')
ax[0].plot(BETAS, [get(ff, 0.0, b)['ref_eta_beta'] for b in BETAS], 's-', color='0.45', ms=3, label='strong-cavity micro-OLED (FWHM 30)')
ax[0].plot(BETAS, [get(ff, 0.0, b)['scat_eta_beta'] for b in BETAS], 'x:', color='0.45', ms=3, label='scatterer only (FWHM 30)')
sh.append(('references_FWHM30', ['beta', 'ideal', 'strong_cavity', 'scatterer_only', 'weak_planar'],
           [[b, get(ff, 0.0, b)['ideal_eta_beta'], get(ff, 0.0, b)['ref_eta_beta'], get(ff, 0.0, b)['scat_eta_beta'], get(ff, 0.0, b)['weakplanar_eta_beta']] for b in BETAS]))
ax[0].set_xlabel(r'acceptance half-angle $\beta$ (deg)'); ax[0].set_ylabel(r'power into cone / generated, $\eta_\beta$')
ax[0].set_title(r'(a) light delivered into the pancake acceptance cone', loc='left'); ax[0].legend(fontsize=5.3); ax[0].grid(alpha=.25)
# (b) gain vs strong-cavity reference, with scatterer loss
for f in FWHMS:
    for a, ls in ((0.0, '-'), (0.02, ':')):
        ax[1].plot(BETAS, [get(f, a, b)['eta_beta'] / get(f, a, b)['ref_eta_beta'] for b in BETAS], ls, marker='o', ms=2.5, color=C[f],
                   label=f'FWHM {f:.0f}' + (' (scatterer loss 2%/pass)' if a else ''))
ax[1].axhline(1, color='k', lw=0.7)
ax[1].set_xlabel(r'$\beta$ (deg)'); ax[1].set_ylabel('system-efficiency gain vs strong cavity')
ax[1].set_title('(b) gain at the eye (same pancake lens)', loc='left'); ax[1].legend(fontsize=5, ncol=2); ax[1].grid(alpha=.25)
sh.append(('gain', ['fwhm', 'a_sc', *[f'beta{b:.0f}' for b in BETAS]],
           [[f, a, *[get(f, a, b)['eta_beta'] / get(f, a, b)['ref_eta_beta'] for b in BETAS]] for f in FWHMS for a in R['a_sc']]))
save(fig, 'figC1_cone_efficiency', sh)

# ---- Fig 2: angular profiles ----
fig, ax = plt.subplots(1, 2, figsize=(7.2, 2.6), gridspec_kw=dict(wspace=0.3))
x = 0.5 * (Mo.AIR_EDGES[:-1] + Mo.AIR_EDGES[1:]); sr = 2 * np.pi * np.sin(np.radians(x)) * np.radians(0.25)
sh = []
for k, b in enumerate((15.0, 25.0)):
    for f in (10.0, 20.0, 50.0):
        r = get(f, 0.0, b); I = np.array(r['hist']) / sr
        ax[k].plot(x, I, color=C[f], lw=1.1, label=f'scatterer+DBR, FWHM {f:.0f}')
        sh.append((f'b{b:.0f}_FWHM{f:.0f}', ['theta_air', 'radiant_intensity'], list(zip(x, I))))
    Ir = np.array(get(20.0, 0.0, b)['ref_hist']) / sr
    ax[k].plot(x, Ir, color='0.4', ls='--', lw=1, label='strong cavity (FWHM 20)')
    sh.append((f'b{b:.0f}_ref', ['theta_air', 'radiant_intensity'], list(zip(x, Ir))))
    ax[k].axvspan(0, b, color='0.92', zorder=0); ax[k].set_xlim(0, 60)
    ax[k].set_xlabel('angle in air (deg)'); ax[k].set_ylabel('radiant intensity (per sr, per generated)')
    ax[k].set_title(f'({"ab"[k]}) designs optimized for beta = {b:.0f} deg', loc='left'); ax[k].legend(fontsize=5.3); ax[k].grid(alpha=.25)
save(fig, 'figC2_angular_profiles', sh)

# ---- Fig 3: blur ----
TF = Mo.fresnel_T(); rows3 = []
fig, ax = plt.subplots(figsize=(3.5, 2.6))
for b, c in ((10, '#2E7D5B'), (20, '#3B6EA5'), (30, '#8E5BA6')):
    Tid = ((Mo.TH_AIR < b) & Mo.ESC).astype(float); qs = [0.85, 0.9, 0.95, 0.97]
    r90 = [blur.spread(Tid, q, n=150000)[1] for q in qs]
    ax.plot(qs, r90, 'o-', color=c, ms=3, label=f'cone {b} deg'); rows3 += [[b, q, v] for q, v in zip(qs, r90)]
r90s = blur.spread(TF, 0.97, n=150000)[1]; ax.axhline(r90s, color='0.5', ls=':', label='scatterer only')
ax.set_xlabel('round-trip survival q = 1 - A\''); ax.set_ylabel(r'$r_{90}$ / gap thickness t')
ax.set_title('(c) lateral spread enclosing 90% of output', loc='left'); ax.legend(fontsize=5.5); ax.grid(alpha=.25)
save(fig, 'figC3_blur', [('r90', ['cone_deg', 'q', 'r90_over_t'], rows3 + [['scatterer only', 0.97, r90s]])])

# ---- impact table ----
print('\nImpact at fixed luminance to the eye (a_sc = 0.01):')
imp = []
for f in FWHMS:
    for b in BETAS:
        r = get(f, 0.01, b); G = r['eta_beta'] / r['ref_eta_beta']
        imp.append([f, b, G, 1 - 1 / G, G ** AF[0], G ** AF[1], PANCAKE * r['eta_beta'], PANCAKE * r['ref_eta_beta']])
        print(f'FWHM {f:4.0f} beta {b:4.0f}: gain x{G:.2f}  panel power -{100 * (1 - 1 / G):.0f}%  lifetime x{G ** AF[0]:.2f}-{G ** AF[1]:.2f}  '
              f'eye/generated {PANCAKE * r["eta_beta"] * 100:.1f}% vs {PANCAKE * r["ref_eta_beta"] * 100:.1f}%')
wb = Workbook(); ws = wb.active; ws.title = 'impact'
ws.append(['fwhm', 'beta', 'gain', 'power_saving', 'lifetime_x_n1.5', 'lifetime_x_n1.6', 'eye_per_generated_new', 'eye_per_generated_ref'])
for d in imp: ws.append([float(v) for v in d])
wb.save('tableC_impact.xlsx')
