import csv, openpyxl
from openpyxl.styles import Font
B = Font(bold=True)
rows = list(csv.DictReader(open('/home/user/OLED-/sim/pass_resolved/pass_resolved.csv')))
C = ['Ag n_sub 1.5', 'Ag n_sub 1.8', 'Al n_sub 1.5', 'Al n_sub 1.8']; K = 50
get = lambda c, f: [float(x[f]) for x in rows if x['case'] == c][:K + 1]
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'README'
for i, l in enumerate([
 'Supplementary Fig. — pass-resolved escape probability p_k, round-trip loss A\'_k and cumulative eta_ext',
 'Stack of Fig. 2(d)-(f): substrate / ITO 50 nm / organics 420 nm (n = 1.8, dipole centred in 20 nm EML) / reflector 100 nm; 550 nm; isotropic; hemispherical MLA with n_MLA = n_sub (LightTools BSDF).',
 'p_k = sum P_sub^(k) B_T / sum P_sub^(k)  (eq. S6);  A\'_k = loss weighted by the light returned at pass k (eq. S7);  eta_ext = sum_k p_k prod_{j<k}(1-p_j)(1-A\'_j)  (eq. S8).',
 'Sheets a_pk, b_Ak, c_cumulative: k = 0-50 (the preview plots 0-15). Sheet reference_lines: p and A\' of eq. (3) (cos.sin weight) and eq. (3) eta_ext = dotted lines. Sheet full: all columns, k = 0-299.',
 'Series converged: sum over k = 0-299 equals the matrix-series eta_ext to 1e-4 (Ag 1.5: 0.9113; Ag 1.8: 0.8755; Al 1.5: 0.7820; Al 1.8: 0.6595).',
 'Source: sim/pass_resolved/pass_resolved.py -> pass_resolved.csv'], 1):
    ws.cell(i, 1, l)
ws.column_dimensions['A'].width = 150
for name, f, sc in (('a_pk', 'p_k', 1), ('b_Ak', "A'_k", 100), ('c_cumulative', 'cumulative eta_ext', 1)):
    s = wb.create_sheet(name); s.cell(1, 1, 'k').font = B
    for j, c in enumerate(C, 2): s.cell(1, j, c + (' (%)' if sc == 100 else '')).font = B
    for k in range(K + 1):
        s.cell(k + 2, 1, k)
        for j, c in enumerate(C, 2): s.cell(k + 2, j, round(get(c, f)[k] * sc, 6))
s = wb.create_sheet('reference_lines')
for j, h in enumerate(['case', 'p (cos.sin)', "A' (cos.sin) (%)", 'eta_ext eq. (3)', 'eta_ext series'], 1): s.cell(1, j, h).font = B
for i, c in enumerate(C, 2):
    r = [x for x in rows if x['case'] == c]
    for j, v in enumerate([c, float(r[0]['p (cos.sin)']), 100 * float(r[0]["A' (cos.sin)"]), float(r[0]['eq3']), float(r[-1]['cumulative eta_ext'])], 1):
        s.cell(i, j, v if j == 1 else round(v, 6))
s = wb.create_sheet('full'); h = list(rows[0].keys())
for j, x in enumerate(h, 1): s.cell(1, j, x).font = B
for i, r in enumerate(rows, 2):
    for j, x in enumerate(h, 1): s.cell(i, j, r[x] if j == 1 else float(r[x]))
wb.save('/home/user/OLED-/sim/pass_resolved/pass_resolved_rawdata.xlsx'); print('ok')
