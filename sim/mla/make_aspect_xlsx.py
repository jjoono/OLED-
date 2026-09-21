# -*- coding: utf-8 -*-
"""Raw data behind the microlens aspect-ratio panel.

Sheet 'aspect sweep' is exactly what the panel plots; the rest is supporting.
"""
import numpy as np, openpyxl, os
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
BOLD = Font(bold=True)

A = np.genfromtxt(os.path.join(HERE, 'mla_aspect.csv'), delimiter=',', names=True,
                  dtype=None, encoding='utf-8')
ALT = np.genfromtxt(os.path.join(HERE, 'mla_aspect_alt.csv'), delimiter=',',
                    names=True, dtype=None, encoding='utf-8')
F = np.genfromtxt(os.path.join(HERE, 'mla_flat_reference.csv'), delimiter=',',
                  names=True, dtype=None, encoding='utf-8')
Z = np.load(os.path.join(HERE, 'mla_bsdf.npz'))
FLAT = {str(r['reflector']): float(r['EQE_flat']) for r in F}
m = A[A['aspect_ratio'] >= 0.09]
NRAY = int(Z['n_ray']) if 'n_ray' in Z else 0

best = {}
for k in ('EQE_Al', 'EQE_Ag'):
    i = int(np.argmax(m[k]))
    w = m['aspect_ratio'][m[k] >= 0.95 * m[k][i]]
    best[k] = (m['aspect_ratio'][i], m[k][i], w[0], w[-1])

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'README'
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 112
rows = [
 ('Microlens aspect ratio - how tight is the tolerance?', None),
 (None, None),
 ('Stack (generic, not the measured device)', None),
 ('  rear', 'Ag 100 nm (McPeak, 0.04382 + 3.81898i at 550 nm) or Al 100 nm (0.95834 + 6.68678i)'),
 ('  ETL', '1.80, 200 nm'),
 ('  EML', '1.80, 20 nm, dipole at the centre, PLQY = 1, isotropic orientation'),
 ('  HTL', '1.80, 200 nm'),
 ('  bottom electrode', 'ITO 50 nm, 1.86362 + 0.00323i (Koenig 2014)'),
 ('  substrate', 'semi-infinite and incoherent, n = 1.80; infinite planar substrate'),
 ('  lens array', 'hexagonally close-packed, index-matched to the substrate (n_MLA = n_sub = 1.80),'),
 (None, 'fill factor pi/(2 sqrt 3) = 0.9069, infinite array, no absorption in the lens'),
 ('  wavelength', '550 nm'),
 (None, None),
 ('Lens shape - and why it changes at AR = 1', None),
 ('  AR <= 1', 'spherical cap of base radius r and height h.  Its sphere has R = (r^2+h^2)/2h and'),
 (None, 'centre z = h - R < 0, i.e. below the base plane, so the widest circle of the cap is'),
 (None, 'the base circle and neighbouring caps touch without ever overlapping.'),
 ('  AR = 1', 'the cap is exactly a hemisphere (R = r, centre on the base plane).'),
 ('  AR > 1', 'the centre would rise above the base plane and the widest circle becomes R > r, so'),
 (None, 'close-packed caps would cut into each other (by 0.0045 r at AR = 1.1, 0.083 r at 1.5).'),
 (None, 'A taller lens is therefore modelled as a half-ellipsoid, (x^2+y^2)/r^2 + z^2/h^2 = 1,'),
 (None, 'whose widest circle is the base circle at every height.  It is the same hemisphere'),
 (None, 'as the cap at AR = 1, so the curve is continuous there by construction.'),
 ('  sheet alt shapes', 'each family outside its own range, for the record.  Forcing the cap above AR = 1'),
 (None, 'puts a spurious dip of about 3 %p near AR = 1.1 - the tracer then meets an internal'),
 (None, 'surface of the overlap region and reads it as an exit.  That dip is a geometry'),
 (None, 'artefact, not physics.  The half-ellipsoid below AR = 1 is a legitimate shape but'),
 (None, 'not a reflowed one: its rim always meets the substrate vertically.'),
 (None, None),
 ('Method', None),
 ('  dipole model', 'uniaxial dipole transfer matrix (CPS); eta_sub is the air + substrate-confined'),
 (None, 'fraction at the first pass, R_LED(theta) the stack reflectance seen from the'),
 (None, 'substrate, P_sub(theta) the angular distribution launched into the substrate.'),
 ('  array model', 'Monte-Carlo ray trace of the lens array, %d rays per 1-degree incidence bin,'
                   % NRAY),
 (None, 'unpolarised Fresnel sampling, periodic boundary, lens-to-lens re-entry, and the'),
 (None, '9.3 %% of the base plane the close-packed circles leave uncovered treated as flat'),
 (None, 'glass/air.  It returns B_T(theta) (escapes to air) and B_R(theta_out, theta_in)'),
 (None, '(returned to the substrate).  B_T + sum_out B_R = 1 to 1e-6 on every column:'),
 (None, 'the array neither absorbs nor traps.'),
 ('  recycling', 'eta_ext = sum_j B_T . v_j with v_1 = P_sub and v_(j+1) = diag(R_LED) B_R v_j,'),
 (None, 'summed to 80 terms - the matrix form, Eq. (1).  The closed form'),
 (None, "eta_ext = p/[p+(1-p)A'] is NOT used: it assumes the surface randomises the angle at"),
 (None, 'every bounce, which a shallow lens does not, and that is exactly what the'),
 (None, 'low-aspect-ratio end of this sweep is about.'),
 ('  EQE', 'eta_sub x eta_ext, at PLQY = 1 and unity charge balance.'),
 (None, None),
 ('Flat reference (dashed lines)', None),
 (None, 'the same stack with no lenses: B_T is the unpolarised Fresnel transmittance and'),
 (None, 'B_R is diagonal, since a flat interface returns light at the same polar angle.'),
 (None, 'EQE = %.4f (Al) and %.4f (Ag).  It cannot be read off the AR -> 0 end of the'
        % (FLAT['Al'], FLAT['Ag'])),
 (None, 'sweep: even a 2 %-aspect-ratio cap still redirects light (EQE 0.227 / 0.310).'),
 (None, None),
 ('What the panel shows', None),
 (None, 'Ag: best EQE %.4f at AR = %.2f, within 5 %% of it from AR = %.2f to %.2f.'
        % (best['EQE_Ag'][1], best['EQE_Ag'][0], best['EQE_Ag'][2], best['EQE_Ag'][3])),
 (None, 'Al: best EQE %.4f at AR = %.2f, within 5 %% of it from AR = %.2f to %.2f.'
        % (best['EQE_Al'][1], best['EQE_Al'][0], best['EQE_Al'][2], best['EQE_Al'][3])),
 (None, 'The optimum is shallow and broad, so the aspect ratio is not a tight tolerance;'),
 (None, 'what matters is that the lenses are there at all and that they close-pack.'),
 (None, None),
 ('Caveats', 'infinite planar substrate and infinite array - no edge, no finite-panel waveguiding.'),
 (None, 'Single wavelength, isotropic emitter, PLQY = 1, lossless index-matched lenses.'),
 (None, 'The numbers are for this generic stack; the shape of the curve is the message.'),
 (None, None),
 ('Sheets', None),
 ('  aspect sweep', 'what the panel plots: p, eta_ext and EQE against AR for both reflectors'),
 ('  flat reference', 'the lens-free interface, same stack'),
 ('  alt shapes', 'each lens family outside its own range (see above)'),
 ('  B_T vs angle', 'escape probability per 1-degree incidence bin, for every AR'),
 ('  B_R at AR=0.50', 'the full return matrix at the Ag optimum, rows = outgoing bin'),
]
for i, (a, b) in enumerate(rows, 1):
    ws.cell(i, 1, a).font = BOLD if (a and not a.startswith(' ')) else Font()
    ws.cell(i, 2, b)


def write(name, arr, fmt=None):
    s = wb.create_sheet(name)
    for j, h in enumerate(arr.dtype.names, 1):
        s.cell(1, j, h).font = BOLD
        s.column_dimensions[get_column_letter(j)].width = max(13, len(h) + 2)
    for i, r in enumerate(arr, 2):
        for j, h in enumerate(arr.dtype.names, 1):
            v = r[h]
            c = s.cell(i, j, v.item() if hasattr(v, 'item') else v)
            if not isinstance(c.value, str):
                c.number_format = (fmt or {}).get(h, '0.000000')
    s.freeze_panes = 'A2'
    return len(arr)


n1 = write('aspect sweep', A, {'aspect_ratio': '0.00'})
n2 = write('flat reference', F)
n3 = write('alt shapes', ALT, {'aspect_ratio': '0.00'})

ar = np.asarray(Z['aspect'], float)
o = np.argsort(ar)
s = wb.create_sheet('B_T vs angle')
s.cell(1, 1, 'theta_in_deg').font = BOLD
s.column_dimensions['A'].width = 13
for j, k in enumerate(o, 2):
    s.cell(1, j, 'AR=%.2f' % ar[k]).font = BOLD
    s.column_dimensions[get_column_letter(j)].width = 10
for i, th in enumerate(Z['theta_deg'], 2):
    s.cell(i, 1, float(th)).number_format = '0.0'
    for j, k in enumerate(o, 2):
        s.cell(i, j, float(Z['BT'][k][i - 2])).number_format = '0.0000'
s.freeze_panes = 'B2'

kk = int(np.argmin(np.abs(ar - 0.50)))
s = wb.create_sheet('B_R at AR=%.2f' % ar[kk])
s.cell(1, 1, 'theta_out \\ theta_in').font = BOLD
s.column_dimensions['A'].width = 18
for j, th in enumerate(Z['theta_deg'], 2):
    s.cell(1, j, float(th)).number_format = '0.0'
for i, th in enumerate(Z['theta_deg'], 2):
    s.cell(i, 1, float(th)).number_format = '0.0'
    for j in range(90):
        s.cell(i, j + 2, float(Z['BR'][kk][i - 2, j])).number_format = '0.00000'
s.freeze_panes = 'B2'

wb.save(os.path.join(HERE, 'mla_aspect_rawdata.xlsx'))
print('mla_aspect_rawdata.xlsx : sweep %d, flat %d, alt %d rows' % (n1, n2, n3))
