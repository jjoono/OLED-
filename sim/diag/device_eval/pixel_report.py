# -*- coding: utf-8 -*-
"""Performance summary (xlsx + png) for one OLED pixel, using OLED-evaluation's own functions.

Differences from the GUI (documented in the README sheet):
  1. angular correction factors = mean of the 0..+90 and -90..0 halves (full -90..90)
  2. calculate_v_correction's negative-angle branch lacks abs(); its sign is fixed here
Usage: python pixel_report.py <jvl.csv> <gon-spec-deg.csv> <out_dir> <name>
"""
import json, os, re, sys
import numpy as np, pandas as pd

EVAL = os.environ.get("OLED_EVAL_DIR", os.path.join(os.path.dirname(os.path.abspath(__file__)), "oled_evaluation"))
sys.path.insert(0, os.path.join(EVAL, "src"))
import evaluation_functions as ef

CAL = os.path.join(EVAL, "calibration")
SPEC_CAL = "Calibration_HDX_QP600_F950SMA-A_2025-09-26.txt"
PIXEL_MM2, PD_DIST_MM, PD_GAIN, PD_AREA_MM2 = 4.0, 58.0, 50, 75.0


def evaluate(jvl_path, deg_path):
    photopic, pd_resp, cie_ref, spec_calib, wl_range = ef.read_calibration_files(
        os.path.join(CAL, "CIE_sle_photopic.csv"), os.path.join(CAL, "Responsivity_PD-extended-range.csv"),
        os.path.join(CAL, "CIE_xyz_1931_2deg.csv"), os.path.join(CAL, SPEC_CAL))
    gain = json.load(open(os.path.join(EVAL, "usr", "photodiode_gain.json")))[str(PD_GAIN)][0]
    pd_res, pd_cut = float(gain["pd_resistance"]), float(gain["pd_cutoff"])

    raw = pd.read_csv(deg_path, sep="\t", skiprows=4)
    raw.columns = [c.strip() for c in raw.columns]
    raw = raw.loc[(raw.wavelength > wl_range[0]) & (raw.wavelength < wl_range[1])]
    ar = raw.drop(columns=[c for c in raw.columns if "_deg" in c])
    cal_ar = ef.calibrate_spectrum(ef.interpolate_spectrum(ar, photopic), spec_calib)
    angs = [c for c in cal_ar.columns if c != "wavelength"]
    pos = cal_ar[[c for c in angs if float(c) >= 0] + ["wavelength"]]   # functions expect wavelength last
    neg = cal_ar[[c for c in angs if float(c) <= 0] + ["wavelength"]]
    c = dict(e_pos=ef.calculate_e_correction(pos), e_neg=ef.calculate_e_correction(neg),
             v_pos=ef.calculate_v_correction(pos, photopic), v_neg=abs(ef.calculate_v_correction(neg, photopic)))
    c["e"], c["v"] = (c["e_pos"] + c["e_neg"]) / 2, (c["v_pos"] + c["v_neg"]) / 2

    perp = pd.DataFrame({"wavelength": raw.wavelength.values, "background": raw["0.0_bg"].values,
                         "intensity": raw["0.0"].values})
    perp = ef.calibrate_spectrum(ef.interpolate_spectrum(perp, photopic), spec_calib)

    j = pd.read_csv(jvl_path, sep="\t", skiprows=7, header=None).apply(pd.to_numeric, errors="coerce")
    j = j.dropna(how="all", axis=1)
    jvl = {"voltage": j[0].values, "current": j[1].values, "pd_voltage": j[2].values}
    sigma = j[3].values if j.shape[1] > 3 else None

    def mk(ec, vc):
        return ef.JVLData(jvl_data=jvl, perpendicular_spectrum=perp, photopic_response=photopic,
                          pd_responsivity=pd_resp, cie_reference=cie_ref, angle_resolved=True,
                          pixel_area=PIXEL_MM2 * 1e-6, pd_resistance=pd_res,
                          pd_radius=np.sqrt(PD_AREA_MM2 * 1e-6 / np.pi), pd_distance=PD_DIST_MM * 1e-3,
                          pd_cutoff=pd_cut, correction_factor=[ec, vc])
    return dict(full=mk(c["e"], c["v"]), gui=mk(c["e_pos"], c["v_pos"]), corr=c, perp=perp,
                cal_ar=cal_ar, photopic=photopic, cie_ref=cie_ref, sigma=sigma)


def build(jvl_path, deg_path, out_dir, name, tag=None):
    os.makedirs(out_dir, exist_ok=True)
    r = evaluate(jvl_path, deg_path); full, gui, c = r["full"], r["gui"], r["corr"]
    df = pd.DataFrame({"Voltage (V)": full.voltage, "Current (mA)": full.current * 1e3,
                       "J (mA/cm2)": full.current_density, "PD voltage (V)": full.pd_voltage,
                       "Luminance (cd/m2)": full.luminance, "EQE (%)": full.eqe,
                       "CE (cd/A)": full.current_efficiency, "PE (lm/W)": full.luminous_efficacy,
                       "Power density (mW/mm2)": full.power_density})
    if r["sigma"] is not None:
        with np.errstate(all="ignore"):
            rel = np.where(full.pd_voltage > 0, r["sigma"] / full.pd_voltage, np.nan)
        df["PD sigma (V)"] = r["sigma"]; df["EQE 1σ (%)"] = full.eqe * rel
    df["EQE 0~+90 only (%)"] = gui.eqe; df["PE 0~+90 only (lm/W)"] = gui.luminous_efficacy
    lit = df[df["Luminance (cd/m2)"] > 0].copy()

    def at_L(col, Lt):
        L = lit["Luminance (cd/m2)"].values
        if Lt < L.min() or Lt > L.max(): return np.nan
        return float(np.interp(np.log10(Lt), np.log10(L), lit[col].values))
    def at_max(col):
        k = lit[col].idxmax(); return lit.loc[k, col], lit.loc[k, "Luminance (cd/m2)"], lit.loc[k, "Voltage (V)"]

    perp = r["perp"]; wl = perp.wavelength.values; I0 = perp.intensity.values
    I0n = I0 / I0.max(); pk = wl[np.argmax(I0)]; above = wl[I0n >= 0.5]; fwhm = above.max() - above.min()
    cal, photopic, cie_ref = r["cal_ar"], r["photopic"], r["cie_ref"]
    xyz = cie_ref[["x_cie", "y_cie", "z_cie"]].values
    rows = []
    for a in sorted([x for x in cal.columns if x != "wavelength"], key=float):
        s = cal[a].values; X, Y, Z = (s[:, None] * xyz).sum(0)
        rows.append(dict(angle=float(a),
                         photon_rel=(s * cal.wavelength.values).sum() / (cal["0.0"].values * cal.wavelength.values).sum(),
                         lum_rel=(s * photopic.photopic_response.values).sum() / (cal["0.0"].values * photopic.photopic_response.values).sum(),
                         lambertian=np.cos(np.deg2rad(float(a))), CIE_x=X / (X + Y + Z), CIE_y=Y / (X + Y + Z),
                         peak_nm=cal.wavelength.values[np.argmax(s)]))
    ang = pd.DataFrame(rows)

    gon_raw_path = deg_path.replace("_gon-spec-deg.csv", "_gon-spec.csv")
    raw = pd.read_csv(gon_raw_path, sep="\t", skiprows=4); raw.columns = [x.strip() for x in raw.columns]
    def area(cn): return ((raw[cn] - raw[cn + "_bg"])[(raw.wavelength > 380) & (raw.wavelength < 780)]).sum()
    deg1, deg2 = area("0.0_deg1") / area("0.0"), area("0.0_deg2") / area("0.0")
    gj = pd.read_csv(deg_path.replace("_gon-spec-deg.csv", "_gon-jvl.csv"), sep="\t", skiprows=6, header=None)
    head = open(deg_path.replace("_gon-spec-deg.csv", "_gon-jvl.csv"), encoding="utf-8", errors="replace").read(200)
    i_set = float(re.search(r"Source Voltage:\s*([0-9.]+)", head).group(1))
    gon_note = f"정전류 {i_set*1e3:.0f} µA 설정, 각도별 전압 {gj[1].min():.3f}~{gj[1].max():.3f} V"
    if gj[1].max() > 6.99:
        gon_note += " (7 V 전압 한계에 걸림: 실제 전류는 설정보다 작았을 수 있음)"

    first = lit.iloc[0]; prev = df[df["Voltage (V)"] < first["Voltage (V)"]]["Voltage (V)"]
    ton = (f"{prev.max():.1f}~{first['Voltage (V)']:.1f}" if first["Luminance (cd/m2)"] > 1 and len(prev)
           else f"{np.interp(0, np.log10(lit['Luminance (cd/m2)']), lit['Voltage (V)']):.2f}")
    S = []
    def add(k, v, u, note=""): S.append((k, v, u, note))
    emax, eL, eV = at_max("EQE (%)"); cmax, cL, cV = at_max("CE (cd/A)"); pmax, pL, pV = at_max("PE (lm/W)")
    add("최대 EQE", emax, "%", f"{eL:.0f} cd/m², {eV:.1f} V")
    add("EQE @ 100 cd/m²", at_L("EQE (%)", 100), "%"); add("EQE @ 1000 cd/m²", at_L("EQE (%)", 1000), "%")
    add("최대 CE", cmax, "cd/A", f"{cL:.0f} cd/m², {cV:.1f} V")
    add("CE @ 100 cd/m²", at_L("CE (cd/A)", 100), "cd/A"); add("CE @ 1000 cd/m²", at_L("CE (cd/A)", 1000), "cd/A")
    add("최대 PE", pmax, "lm/W", f"{pL:.1f} cd/m², {pV:.1f} V")
    add("PE @ 100 cd/m²", at_L("PE (lm/W)", 100), "lm/W"); add("PE @ 1000 cd/m²", at_L("PE (lm/W)", 1000), "lm/W")
    add("전압 @ 100 cd/m²", at_L("Voltage (V)", 100), "V"); add("전압 @ 1000 cd/m²", at_L("Voltage (V)", 1000), "V")
    add("turn-on 전압 (1 cd/m²)", ton, "V", f"첫 발광점 {first['Voltage (V)']:.1f} V에서 {first['Luminance (cd/m2)']:.1f} cd/m²")
    add("최대 휘도", lit["Luminance (cd/m2)"].max(), "cd/m²",
        f"{lit['Voltage (V)'].iloc[-1]:.1f} V, J {lit['J (mA/cm2)'].iloc[-1]:.2f} mA/cm²")
    add("CIE (x, y), 0°", f"({full.cie_coordinates[0]:.3f}, {full.cie_coordinates[1]:.3f})", "")
    add("피크 파장, 0°", pk, "nm"); add("FWHM, 0°", fwhm, "nm")
    add("각도 보정 계수 e (-90~90 평균)", c["e"], "", f"0~+90: {c['e_pos']:.4f}, -90~0: {c['e_neg']:.4f}, Lambertian = 0.5")
    add("각도 보정 계수 v (-90~90 평균)", c["v"], "", f"0~+90: {c['v_pos']:.4f}, -90~0: {c['v_neg']:.4f}")
    add("EQE 차이: 0~+90만 쓴 GUI 방식 대비", 100 * (c["e"] / c["e_pos"] - 1), "%")
    add("각도 측정 중 0° 밝기 변화 (deg1, deg2)", f"{deg1:.3f}, {deg2:.3f}", "× 처음 0°", "원본 gon-spec 기준, 계산에는 deg 파일 사용")
    add("각도 측정 구동 조건", "", "", gon_note)
    summ = pd.DataFrame(S, columns=["항목", "값", "단위", "비고"])

    # ---- xlsx
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter as L_
    Fn = Font(name="Arial", size=10); Fb = Font(name="Arial", size=10, bold=True)
    HEAD = PatternFill("solid", start_color="D9E1F2")
    wb = Workbook(); ws = wb.active; ws.title = "README"
    notes = [f"{name} — 소자 성능 정리", "",
             f"JVL 파일: {jvl_path}", f"각도 스펙트럼: {deg_path}",
             "계산: OLED-evaluation(src, 2025-02 사본)의 함수 그대로 사용 (JVLData, calibrate_spectrum, calculate_e/v_correction)",
             f"조건: 픽셀 {PIXEL_MM2:g} mm², PD–소자 {PD_DIST_MM:g} mm, PD {PD_GAIN} dB, PD 면적 {PD_AREA_MM2:g} mm², pd_cutoff 1e-4 V, 분광기 보정 {SPEC_CAL}",
             f"각도 분포: gon-spec-deg (열화 보정본), {gon_note}, 0° 스펙트럼을 수직 스펙트럼으로 사용",
             "GUI와 다른 점 1: 각도 보정 계수를 0~+90°와 -90~0°에서 각각 같은 함수로 계산해 평균 (GUI는 0~+90°만 사용)",
             "GUI와 다른 점 2: calculate_v_correction 의 음의 각도 분기에 abs()가 빠진 버그를 절댓값으로 보정 (PE에만 영향)",
             "PE = luminous efficacy (lm/W). EQE·PE·전력밀도는 각도 분포 보정, 휘도·CE는 정면(0°) 기준",
             "EQE 1σ: Accurate 모드의 PD sigma ÷ PD 로 계산한 통계 오차 (있을 때만)", "",
             "주의: 절대 EQE는 PD 거리 기준점에 민감함 (거리 ±1 mm → EQE 약 ±3.5%)"]
    for i, t in enumerate(notes, 1): ws.cell(row=i, column=1, value=t).font = Fb if i == 1 else Fn
    ws.column_dimensions["A"].width = 140

    def sheet(nm, frame, w=14, idx=None):
        s = wb.create_sheet(nm) if idx is None else wb.create_sheet(nm, idx)
        for jj, h in enumerate(frame.columns, 1):
            cc = s.cell(row=1, column=jj, value=h); cc.font = Fb; cc.fill = HEAD; s.column_dimensions[L_(jj)].width = w
        for ii, row in enumerate(frame.itertuples(index=False), 2):
            for jj, v in enumerate(row, 1):
                if isinstance(v, (np.floating, np.integer)): v = v.item()
                if isinstance(v, float) and not np.isfinite(v): v = None
                s.cell(row=ii, column=jj, value=v).font = Fn
        s.freeze_panes = "A2"; return s
    sheet("JVL", df.round(6))
    sheet("Angular", ang.rename(columns={"angle": "Angle (deg)", "photon_rel": "Photon intensity / 0°",
                                         "lum_rel": "Luminous intensity / 0°", "lambertian": "Lambertian cos(θ)"}).round(5), 18)
    spec = pd.DataFrame({"Wavelength (nm)": wl, "0° normalized": I0n})
    sheet("Spectrum", spec[(spec["Wavelength (nm)"] >= 380) & (spec["Wavelength (nm)"] <= 780)].round(5), 16)
    ss = wb.create_sheet("Summary", 1)
    for jj, h in enumerate(["항목", "값", "단위", "비고"], 1):
        cc = ss.cell(row=1, column=jj, value=h); cc.font = Fb; cc.fill = HEAD
    n = len(df) + 1; col = {h: L_(i + 1) for i, h in enumerate(df.columns)}
    formula = {k: f"=MAX(JVL!{col[h]}2:{col[h]}{n})" for k, h in
               [("최대 EQE", "EQE (%)"), ("최대 CE", "CE (cd/A)"), ("최대 PE", "PE (lm/W)"), ("최대 휘도", "Luminance (cd/m2)")]}
    for ii, (k, v, u, note) in enumerate(S, 2):
        ss.cell(row=ii, column=1, value=k).font = Fn
        val = formula.get(k, v.item() if isinstance(v, np.floating) else v)
        if isinstance(val, float): val = None if not np.isfinite(val) else round(val, 4)
        cv = ss.cell(row=ii, column=2, value=val); cv.font = Fn
        if isinstance(v, (float, np.floating)): cv.number_format = "0.00"
        ss.cell(row=ii, column=3, value=u).font = Fn
        ss.cell(row=ii, column=4, value=note + ("  (값: JVL 시트 로그휘도 선형보간)" if "@" in k else "")).font = Fn
    for jj, w in enumerate([36, 14, 12, 80], 1): ss.column_dimensions[L_(jj)].width = w
    tag = tag or re.sub(r"[^\w]+", "_", name.split()[-1]).strip("_")
    xl = os.path.join(out_dir, f"{tag}_performance.xlsx"); wb.save(xl)

    # ---- figure
    import matplotlib; matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from matplotlib.ticker import FuncFormatter
    plt.rcParams.update({"font.family": ["Malgun Gothic", "NanumGothic", "Noto Sans CJK KR", "DejaVu Sans"], "font.size": 9, "axes.unicode_minus": False,
                         "mathtext.fontset": "dejavusans", "axes.spines.top": False, "axes.spines.right": False})
    B, G, GR = "#2a78d6", "#1baf7a", "#9a9890"
    fig, ax = plt.subplots(2, 4, figsize=(15, 7.2)); a = ax.ravel()
    on = df[df["Voltage (V)"] >= 2.0]
    a[0].semilogy(on["Voltage (V)"], on["J (mA/cm2)"].clip(lower=1e-5), "o-", color=B, ms=3, lw=1.6)
    a[0].set(xlabel="Voltage (V)", ylabel="J (mA/cm²)", title="J–V")
    a[1].semilogy(lit["Voltage (V)"], lit["Luminance (cd/m2)"], "o-", color=B, ms=3, lw=1.6)
    a[1].set(xlabel="Voltage (V)", ylabel="Luminance (cd/m²)", title="L–V")
    for k, cn in enumerate(["EQE (%)", "CE (cd/A)", "PE (lm/W)"]):
        A = a[2 + k]
        A.semilogx(lit["Luminance (cd/m2)"], lit[cn], "o-", color=B, ms=3, lw=1.6, label="-90~90° 보정")
        if cn == "EQE (%)":
            if "EQE 1σ (%)" in lit:
                A.errorbar(lit["Luminance (cd/m2)"], lit[cn], yerr=lit["EQE 1σ (%)"], fmt="none", ecolor=B, elinewidth=0.8, capsize=0)
            A.semilogx(lit["Luminance (cd/m2)"], lit["EQE 0~+90 only (%)"], "--", color=GR, lw=1.2, label="0~+90°만 (GUI)")
            A.legend(frameon=False, fontsize=8, loc="lower left")
        A.set(xlabel="Luminance (cd/m²)", ylabel=cn, title=f"{cn.split(' ')[0]}–L", ylim=(0, lit[cn].max() * 1.15))
    a[5].plot(ang.angle, ang.photon_rel, "o-", color=B, ms=3.5, lw=1.6, label="측정 (광자)")
    a[5].plot(ang.angle, ang.lambertian, "--", color=GR, lw=1.2, label="Lambertian")
    a[5].set(xlabel="Angle (°)", ylabel="Intensity / 0°", title=f"각도 분포 ({i_set*1e3:.0f} µA)", xlim=(-92, 92), ylim=(0, 1.1))
    a[5].set_xticks(range(-90, 91, 30)); a[5].legend(frameon=False, fontsize=8, loc="lower center")
    m = (wl >= 400) & (wl <= 700)
    a[6].plot(wl[m], I0n[m], color=G, lw=1.8)
    a[6].set(xlabel="Wavelength (nm)", ylabel="Normalized intensity", title=f"0° 스펙트럼 (피크 {pk:.0f} nm)", ylim=(0, 1.05))
    a[7].semilogx(lit["J (mA/cm2)"], lit["EQE (%)"], "o-", color=B, ms=3, lw=1.6)
    a[7].set(xlabel="J (mA/cm²)", ylabel="EQE (%)", title="EQE–J", ylim=(0, lit["EQE (%)"].max() * 1.15))
    dec = FuncFormatter(lambda v, _: f"$10^{{{int(round(np.log10(v)))}}}$" if v > 0 and abs(np.log10(v) - round(np.log10(v))) < 1e-6 else "")
    for A in a:
        A.grid(alpha=0.25, lw=0.6)
        if A.get_xscale() == "log": A.xaxis.set_major_formatter(dec)
        if A.get_yscale() == "log": A.yaxis.set_major_formatter(dec)
    fig.suptitle(f"{name}  ·  {PIXEL_MM2:g} mm², PD {PD_DIST_MM:g} mm, {PD_GAIN} dB, gon-spec-deg (-90~90°)", fontsize=11)
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    png = os.path.join(out_dir, f"{tag}_performance.png"); fig.savefig(png, dpi=150); plt.close(fig)
    return summ, xl, png, df


if __name__ == "__main__":
    summ, xl, png, _ = build(*sys.argv[1:5])
    pd.set_option("display.width", 220); pd.set_option("display.max_colwidth", 90)
    print(summ.to_string(index=False)); print("saved:", xl, png)
