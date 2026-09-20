# -*- coding: utf-8 -*-
"""Raw data behind Fig. 3(c)."""
import numpy as np, openpyxl, os
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import nspp, materials as M

HERE = os.path.dirname(os.path.abspath(__file__))
BOLD = Font(bold=True)
wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'README'
ws.column_dimensions['A'].width = 26
ws.column_dimensions['B'].width = 112
rows = [
 ('Fig. 3(c) - a low out-of-plane index only helps past a threshold', None),
 (None, None),
 ('Stack', None),
 ('  top / rear', 'air | Ag 100 nm  (McPeak, 0.04382 + 3.81898i at 550 nm)'),
 ('  ETL', 'uniaxial, n_o and n_e swept, thickness d_ETL swept'),
 ('  EML', '1.80 isotropic, 20 nm, dipole at the centre, PLQY = 1, isotropic orientation'),
 ('  HTL', '1.80 isotropic, 50 nm'),
 ('  bottom electrode', 'ITO 50 nm, 1.86362 + 0.00323i  (Koenig 2014, material.l_ITO)'),
 ('  substrate', 'semi-infinite and incoherent, n = 1.80'),
 ('  wavelength', '550 nm'),
 (None, None),
 ('Method', None),
 ('  model', 'dipole transfer matrix (CPS), uniaxial layers, five channels summing to 1'),
 (None, 'air + substrate-confined + waveguided + SPP/evanescent + absorbed'),
 ('  quadrature', 'u = sin(th) below the light line, u = sqrt(1+v^2) above it, composite Simpson,'),
 (None, '12000 panels per sub-interval; the branch point u = 1 is approached to 1e-5.'),
 (None, 'Reproduces the validated Fig. 3(b) solver to 3e-16 in the isotropic limit and the'),
 (None, "author's TMF_birefringence_whole.m to 6e-9 for a uniaxial stack."),
 ('  eta_sub', 'air + substrate-confined. n_sub = n_EML = 1.80, so there is no waveguided channel:'),
 (None, 'everything below k_x/k0 = 1.80 reaches the substrate and everything above it does not.'),
 (None, None),
 ('n_SPP', None),
 ('  definition', 'TM surface-plasmon index at a metal / uniaxial-dielectric interface,'),
 (None, '(k_SPP/k0)^2 = eps_m (eps_o - eps_m)/(eps_o - eps_m^2/eps_e), the semi-infinite limit.'),
 (None, 'A finite ETL hybridises the plasmon with the isotropic layers above it, so the mode'),
 (None, 'index at d_ETL = 60-150 nm sits above this value; that is why the step in eta_sub'),
 (None, 'occurs at n_SPP = 1.64 (60 nm), 1.73 (100 nm), 1.76 (150 nm) rather than at 1.80.'),
 (None, None),
 ('Measured ETLs at 550 nm (nk_JH_total.mat)', None),
]
for k, (no, ne) in M.ETL.items():
    rows.append(('  ' + k, 'n_o = %.4f   n_e = %.4f   ->   n_SPP(Ag) = %.4f   %s'
                 % (no, ne, nspp.n_spp(M.AG, no, ne),
                    'leaky into n_sub = 1.80' if nspp.n_spp(M.AG, no, ne) < 1.8
                    else 'still bound')))
rows += [(None, None), ('Sheets', None),
         ('  spectrum', 'TM and TE dissipation density per unit k_x/k0, at three ETL thicknesses'),
         ('  eta_sub vs n_SPP', 'three n_o families and the measured ETLs, at three ETL thicknesses'),
         ('  thickness sweep', 'n_o = 1.80, n_e = 1.40...1.80, d_ETL = 10...400 nm'),
         ('  threshold', 'the ETL thickness needed for SPP < 10 % and for eta_sub > 90 %')]
for i, (a, b) in enumerate(rows, 1):
    ws.cell(i, 1, a).font = BOLD if (a and not a.startswith(' ')) else Font()
    ws.cell(i, 2, b)


def sheet(name, path, fmt=None):
    A = np.genfromtxt(os.path.join(HERE, path), delimiter=',', names=True,
                      dtype=None, encoding='utf-8')
    s = wb.create_sheet(name)
    for j, h in enumerate(A.dtype.names, 1):
        s.cell(1, j, h).font = BOLD
        s.column_dimensions[get_column_letter(j)].width = max(11, len(h) + 2)
    for i, r in enumerate(A, 2):
        for j, h in enumerate(A.dtype.names, 1):
            v = r[h]
            c = s.cell(i, j, v.item() if hasattr(v, 'item') else v)
            if not isinstance(c.value, str):
                c.number_format = (fmt or {}).get(h, '0.000000')
    s.freeze_panes = 'A2'
    return len(A)


n1 = sheet('spectrum', 'fig3c_spectrum.csv',
           {'d_ETL_nm': '0', 'n_eff': '0.00000', 'TM': '0.000E+00', 'TE': '0.000E+00'})
n2 = sheet('eta_sub vs n_SPP', 'fig3c_collapse.csv', {'d_ETL_nm': '0.0'})
n3 = sheet('thickness sweep', 'fig3c_sweep.csv', {'d_ETL_nm': '0.0'})
n4 = sheet('threshold', 'fig3c_threshold.csv',
           {'d_SPP_below_10pct_nm': '0.0', 'd_eta_sub_above_90pct_nm': '0.0'})
wb.save(os.path.join(HERE, 'fig3c_rawdata.xlsx'))
print('fig3c_rawdata.xlsx : %d / %d / %d / %d rows' % (n1, n2, n3, n4))
