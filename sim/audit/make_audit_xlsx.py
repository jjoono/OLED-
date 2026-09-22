# -*- coding: utf-8 -*-
"""Workbook of the audit: every closed-form panel recomputed with the series, side by side."""
import numpy as np, openpyxl, os, scipy.io as sio
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__)); BOLD = Font(bold=True)
def load(fn): return np.genfromtxt(os.path.join(HERE, fn), delimiter=',', names=True, dtype=None, encoding='utf-8')
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'README'
ws.column_dimensions['A'].width = 24; ws.column_dimensions['B'].width = 112
rows = [
 ('Data audit — closed form (eq. 3) against the matrix series (eq. 1)', None),
 (None, 'Every panel that had been evaluated with eta_ext = p/[p+(1-p)A\'] is recomputed here with the full series'),
 (None, 'P_out = P0 P_sub (I - B_R R)^-1 B_T^T using the author\'s LightTools hemisphere BSDF (lt_hemisphere_bsdf.mat, slice n_MLA = n_sub).'),
 (None, 'The two agree only when the returned light is angularly randomised; the close-packed hemisphere returns its light at'),
 (None, 'steeper angles, so the series is 3-6 %p below the closed form at n_sub 1.5-1.8 and close to it at 1.9-2.0.'),
 (None, None),
 ('Stack', 'air | reflector 100 nm | ETL 1.8 | EML 1.8 20 nm (dipole at the centre) | HTL 1.8 | TCO | substrate n_sub; 550 nm, PLQY 1, Theta 2/3 unless swept.'),
 (None, 'Ag McPeak entry of the author\'s library, Al Johnson-Christy type, ITO Koenig 2014 (1.864 + 0.0032i).'),
 ('p', 'p_cos_sin: B_T weighted by cos(theta)sin(theta) — a property of the outcoupling structure alone (this is what the text should call p).'),
 (None, 'p_Psub: B_T weighted by the device\'s own P_sub(theta) — depends on the device; given for reference only.'),
 ('tail', 'the last (100th) term of the series over the sum — convergence check, all < 1e-6.'),
 (None, None),
 ('Sheets', None),
 ('  fig2c', 'ITO thickness at n_sub 1.5 (slice 5), Al and Ag: A\', eta_ext closed (p = 0.427) and series'),
 ('  fig2def', 'substrate index 1.30-2.00 (slice n), ITO 50 nm: p, A\', eta_sub, eta_ext and EQE closed and series; old = the closed-form curves as drawn (author\'s p)'),
 ('  fig3a', 'k_ITO (ITO 50 nm) and n_Ag (10 nm Ag, k 3.819) at n_sub 1.8 (slice 11): closed and series'),
 ('  fig5c', 'dipole orientation: reference Al/n 1.5/80+230 nm, proposed Ag/n 1.8/200+200 nm: closed and series'),
 ('  table1', 'real material stack at n_sub 1.8: B3PyMPM ETL, TCTA:B3PyMPM EML, TAPC HTL, ITO 50 nm (k 0.0032 / 0.002), Ag'),
 ('  BSDF LT vs tracer', 'B_T(theta) of the hemisphere at n 1.5 and 1.8 from LightTools and from this repository\'s tracer'),
 ('  p per slice', 'p_cos_sin of all fifteen LightTools slices'),
]
for i, (a, b) in enumerate(rows, 1):
    ws.cell(i, 1, a).font = BOLD if (a and not a.startswith(' ')) else Font(); ws.cell(i, 2, b)
def sheet(name, arr, extra=None):
    s = wb.create_sheet(name); names = list(arr.dtype.names)
    for j, h in enumerate(names, 1):
        s.cell(1, j, h).font = BOLD; s.column_dimensions[get_column_letter(j)].width = max(12, len(h) + 2)
    for i, r in enumerate(arr, 2):
        for j, h in enumerate(names, 1):
            v = r[h]; c = s.cell(i, j, v.item() if hasattr(v, 'item') else v)
            if not isinstance(c.value, str): c.number_format = '0.0000'
    if extra is not None:
        col = len(names) + 2
        for j, h in enumerate(extra.dtype.names, col):
            s.cell(1, j, 'old_' + h).font = BOLD; s.column_dimensions[get_column_letter(j)].width = max(12, len(h) + 6)
        for i, r in enumerate(extra, 2):
            for j, h in enumerate(extra.dtype.names, col):
                v = r[h]; c = s.cell(i, j, v.item() if hasattr(v, 'item') else v)
                if not isinstance(c.value, str): c.number_format = '0.0000'
    s.freeze_panes = 'A2'
sheet('fig2c', load('fig2c_series.csv'))
sheet('fig2def', load('fig2def_series.csv'), np.genfromtxt('/home/user/OLED-/figures/fig2b_mock/fig2b_curves_konig.csv', delimiter=',', names=True, dtype=None, encoding='utf-8'))
sheet('fig3a', load('fig3a_series.csv'))
sheet('fig5c', load('fig5c_series.csv'))
sheet('table1', load('table1_series.csv'))
Z = np.load('/home/user/OLED-/sim/mla/lt_vs_tracer_hemisphere.npz')
s = wb.create_sheet('BSDF LT vs tracer')
for j, h in enumerate(['theta_deg', 'LT n1.5 B_T', 'tracer n1.5 B_T', 'LT n1.8 B_T', 'tracer n1.8 B_T'], 1): s.cell(1, j, h).font = BOLD; s.column_dimensions[get_column_letter(j)].width = 16
for i in range(90):
    for j, v in enumerate([i + 0.5, Z['lt15'][i], Z['mine15'][i], Z['lt18'][i], Z['mine18'][i]], 1): s.cell(2 + i, j, float(v)).number_format = '0.0000'
B = sio.loadmat('/home/user/OLED-/sim/mla/lt_hemisphere_bsdf.mat')['BSDF_MLA']; NS = np.round(np.arange(1.30, 2.001, 0.05), 2)
TH = np.arange(90) + 0.5; W = np.cos(np.radians(TH)) * np.sin(np.radians(TH))
s = wb.create_sheet('p per slice')
for j, h in enumerate(['slice', 'n_MLA', 'p_cos_sin', '1/n^2', 'closure min', 'closure max'], 1): s.cell(1, j, h).font = BOLD; s.column_dimensions[get_column_letter(j)].width = 14
for k, n in enumerate(NS):
    BT = B[:90, :, k].sum(0); cs = B[:, :, k].sum(0)
    for j, v in enumerate([k + 1, float(n), float(np.sum(BT * W) / np.sum(W)), float(1 / n**2), float(cs.min()), float(cs.max())], 1):
        s.cell(2 + k, j, v).number_format = '0.0000' if j > 2 else '0.00'
wb.save(os.path.join(HERE, 'data_audit_recompute.xlsx')); print('saved data_audit_recompute.xlsx')
