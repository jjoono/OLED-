# -*- coding: utf-8 -*-
"""Raw data behind Fig. 3(a): electrode quality, for the two alternative bottom electrodes.
eta_ext and EQE are live formulas driven by a single p cell."""
import numpy as np, openpyxl, os
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
BOLD = Font(bold=True); YEL = PatternFill('solid', fgColor='FFFF00')
F4 = '0.0000'
A = np.loadtxt(os.path.join(HERE, 'f3a_ito_konig.csv'), delimiter=',')
B = np.loadtxt(os.path.join(HERE, 'f3a_ag_konig.csv'), delimiter=',')
# columns: param, eta_sub, A', wg, u>1, abs, abs_top, abs_bottom, ext(.30), ext(.40), EQE(.30), EQE(.40)

wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'README'
ws.column_dimensions['A'].width = 30; ws.column_dimensions['B'].width = 108
rows = [
 ('Fig. 3(a) — electrode quality', None),
 ('What the round-trip loss and the substrate-delivered power do as the transparent electrode gets worse', None),
 (None, None), ('Model', None),
 ('Method', 'Dipole transfer matrix (CPS) with the five-channel power budget: air + substrate-confined + waveguided + u>1 + absorbed = 1'),
 ('Wavelength', '550 nm, single wavelength'),
 ('Emitter', 'isotropic dipole (horizontal fraction 2/3) at the centre of the EML, PLQY = 1'),
 ('u grid', '3000 points, maximum u = 3; the five channels close to 1.000000 on every row'),
 (None, None), ('Stack — two alternative devices, one bottom electrode each', None),
 ('Common part', 'Ag 100 nm (McPeak measured n,k = 0.044 + 3.819i, never swept) / ETL 200 nm / EML 20 nm / HTL 200 nm / [bottom electrode] / substrate'),
 ('Device A (sheet a)', 'bottom electrode = 50 nm TCO, n = 1.8636 + k i, with k swept from 0 to 0.08'),
 ('Device B (sheet b)', 'bottom electrode = 10 nm thin Ag, n = n_Ag + 3.819i, with n_Ag swept from 0 to 0.5'),
 ('Organics', 'ETL, EML and HTL isotropic, n = 1.8, k = 0 — the same generic stack as Fig. 2'),
 ('Substrate', 'n = 1.80, index-matched to the outcoupling structure'),
 ('TCO real index', '1.8636, the real part of the Koenig et al. (2014) ITO at 550 nm, so the sweep passes through the operating point of Fig. 2'),
 (None, None), ('Cross-check against Fig. 2', None),
 ('same stack, same numbers', "at k_TCO = 0.0032 (the Koenig value) device A gives eta_sub = 0.9564 and A' = 0.0310, identical to the Ag / ITO 50 nm point of Fig. 2 at n_sub = 1.8"),
 (None, None), ('Quantities', None),
 ('eta_sub', 'substrate-delivered power: the fraction of dipole power that reaches the substrate at all angles (simulated)'),
 ("A'", "round-trip loss, 1 minus the reflectance of the stack seen from the substrate, flux-weighted over angle (simulated)"),
 ('waveguided, u>1, absorbed', 'the other three channels of the budget.  u>1 is surface plasmon plus any light guided in the electrode when its index exceeds the substrate'),
 ('absorbed, split', 'the same absorption separated into the Ag reflector and the bottom electrode'),
 ('eta_ext', "p / [p + (1 - p) A'].  FORMULA, driven by the p cell below"),
 ('EQE', 'eta_sub x eta_ext.  FORMULA'),
 (None, None), ('Input — edit this cell', None),
 ('p (single-pass escape)', 0.30),
 ('Source of p', 'read off the single-pass escape probability of the outcoupling structure at n_sub = 1.8'),
 (None, None), ('Provenance', None),
 ('Script', 'sim/design_rule4/dr4e.m (Octave), MODE=kito for sheet a and MODE=nagb for sheet b, run with UNUM=3000'),
 ('Source files', 'f3a_ito_konig.csv and f3a_ag_konig.csv'),
 ('Superseded', 'p_ito_n19.csv and p_ag.csv, the earlier run at n_TCO = 1.9 with a uniaxial ETL (n_e = 1.6) and a thin-Ag k of 3.5'),
 ('Repository', 'jjoono/OLED-, branch claude/oled-efficiency-paper-ag-955zsk'),
]
for i, (a, b) in enumerate(rows, start=1):
    if a is not None: ws.cell(row=i, column=1, value=a).font = BOLD
    if b is not None:
        c = ws.cell(row=i, column=2, value=b)
        if isinstance(b, float): c.fill = YEL; c.font = BOLD; c.number_format = '0.00'
ws['A1'].font = Font(bold=True, size=14)
P_ROW = [i for i, (a, _) in enumerate(rows, start=1) if a == 'p (single-pass escape)'][0]

def sheet(title, note, param, D):
    s = wb.create_sheet(title)
    s['A1'] = note; s['A1'].font = BOLD
    s['A2'] = f"eta_ext and EQE are formulas driven by p in README!B{P_ROW}; the closure column re-adds the four channels"
    hdr = [param, 'eta_sub', "A' (round trip)", 'waveguided', 'u>1 (SPP etc.)', 'absorbed',
           '  in the Ag reflector', '  in the bottom electrode', 'eta_ext', 'EQE', 'closure check']
    for j, h in enumerate(hdr, start=1): s.cell(row=4, column=j, value=h).font = BOLD
    P = f'README!$B${P_ROW}'
    for i in range(len(D)):
        r = 5 + i
        for j, v in enumerate(D[i, :8], start=1):
            s.cell(row=r, column=j, value=float(v)).number_format = F4
        s.cell(row=r, column=9, value=f'={P}/({P}+(1-{P})*C{r})').number_format = F4
        s.cell(row=r, column=10, value=f'=B{r}*I{r}').number_format = F4
        s.cell(row=r, column=11, value=f'=B{r}+D{r}+E{r}+F{r}').number_format = '0.000000'
    s.column_dimensions['A'].width = 16
    for j in range(2, 12): s.column_dimensions[get_column_letter(j)].width = 22 if j in (7, 8) else 16
    s.freeze_panes = 'A5'

sheet('a_TCO_electrode', 'Device A: 50 nm TCO bottom electrode, n = 1.8636 + k i', 'k_TCO', A)
sheet('b_thin_Ag_electrode', 'Device B: 10 nm thin Ag bottom electrode, n = n_Ag + 3.819i', 'n_Ag', B)
wb.save(os.path.join(HERE, 'fig3a_rawdata.xlsx'))
print('saved fig3a_rawdata.xlsx')
