"""Raw data behind 'Design rule: parasitic absorption sets the substrate-to-air
extraction efficiency' (dr4e_mock.png) as a workbook.
eta_ext and EQE are written as live formulas off a single p input cell."""
import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = 'Arial'
HDR_FILL = PatternFill('solid', fgColor='1F3864')
HDR_FONT = Font(name=FONT, size=10, bold=True, color='FFFFFF')
IN_FILL = PatternFill('solid', fgColor='FFFF00')
SUB_FONT = Font(name=FONT, size=10, bold=True, color='1F3864')
BODY = Font(name=FONT, size=10)
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

A = np.loadtxt('p_ito_n19.csv', delimiter=',')   # device A, k_TCO sweep
B = np.loadtxt('p_ag.csv', delimiter=',')        # device B, n_Ag sweep
P_CELL = None   # set once the README row that holds p is known

wb = Workbook()

# ---------------------------------------------------------------- README
ws = wb.active; ws.title = 'README'
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 34; ws.column_dimensions['B'].width = 96
rows = [
    ('title', 'Design rule: parasitic absorption sets the substrate-to-air extraction efficiency'),
    ('sub', 'Raw data behind the three panels of dr4e_mock.png'),
    ('blank', ''),
    ('h', 'What this is'),
    ('kv', ('Model', 'Dipole transfer-matrix (CPS) with the five-channel power budget: air + substrate-confined + waveguided + u>1 + absorption = 1')),
    ('kv', ('Wavelength', '550 nm, single wavelength')),
    ('kv', ('Emitter', 'Isotropic dipole (horizontal fraction 2/3) at the centre of the EML; photoluminescence quantum yield = 1')),
    ('kv', ('u grid', '3000 points, maximum u = 3; five-channel closure better than 1e-6')),
    ('blank', ''),
    ('h', 'Stacks — two alternative devices, not one'),
    ('kv', ('Common part', 'Ag 100 nm (McPeak measured n,k = 0.044 + 3.819i, never swept) / ETL 200 nm / EML 20 nm / HTL 200 nm / [bottom electrode] / substrate')),
    ('kv', ('Device A (sheet a)', 'Bottom electrode = transparent conducting oxide, 50 nm, n = 1.9 + k_TCO i;  k_TCO is swept')),
    ('kv', ('Device B (sheet b)', 'Bottom electrode = thin Ag, 10 nm, n = n_Ag + 3.5i;  n_Ag is swept')),
    ('kv', ('Organic indices', 'EML and HTL isotropic at n = 1.8, k = 0. ETL uniaxial: ordinary n_o = 1.8, extraordinary n_e = 1.6')),
    ('kv', ('Substrate', 'n = 1.8')),
    ('blank', ''),
    ('h', 'Quantities'),
    ('kv', ('eta_sub', 'Substrate-delivered power: the fraction of dipole power that reaches the substrate, at all angles (simulated)')),
    ('kv', ("A'", "Round-trip loss, 1 minus the reflectance of the OLED stack seen from the substrate, flux-weighted (cos x sin) over substrate angles and averaged over p and s polarisation (simulated)")),
    ('kv', ('eta_ext', "Substrate-to-air extraction efficiency = p / [p + (1 - p) A'].  FORMULA, recalculates from the p cell below")),
    ('kv', ('EQE', 'eta_sub x eta_ext.  FORMULA')),
    ('kv', ('u>1 bin', 'Power beyond the light line: surface plasmon, plus light guided in the electrode when its index exceeds the substrate (simulated)')),
    ('blank', ''),
    ('h', 'Input — edit this cell and the whole workbook updates'),
    ('in', ('p (single-pass escape probability of the microlens array)', 0.30)),
    ('kv', ('Source of p', 'Read off the blue single-pass-escape-probability curve of Fig. 1c at n_sub = 1.8; digitisation uncertainty about +/- 0.01. The slide used 0.40.')),
    ('blank', ''),
    ('h', 'Provenance'),
    ('kv', ('Simulation script', 'sim/design_rule4/dr4d.m (Octave 8.4), run with UNUM=3000')),
    ('kv', ('Source files', 'p_ito_n19.csv (sheet a), p_ag.csv (sheet b)')),
    ('kv', ('Figure script', 'sim/design_rule4/plot_dr4e.py')),
    ('kv', ('Repository', 'jjoono/OLED-, branch claude/oled-efficiency-paper-ag-955zsk')),
]
r = 1
for kind, val in rows:
    if kind == 'title':
        ws.cell(r, 1, val).font = Font(name=FONT, size=14, bold=True, color='1F3864')
    elif kind == 'sub':
        ws.cell(r, 1, val).font = Font(name=FONT, size=10, italic=True, color='595959')
    elif kind == 'h':
        ws.cell(r, 1, val).font = SUB_FONT
    elif kind == 'kv':
        ws.cell(r, 1, val[0]).font = Font(name=FONT, size=10, bold=True)
        c = ws.cell(r, 2, val[1]); c.font = BODY; c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.row_dimensions[r].height = max(14, 13*(1 + len(str(val[1]))//110))
    elif kind == 'in':
        p_row = r
        ws.cell(r, 1, val[0]).font = Font(name=FONT, size=10, bold=True)
        c = ws.cell(r, 2, val[1]); c.font = Font(name=FONT, size=11, bold=True, color='0000FF')
        c.fill = IN_FILL; c.border = BOX; c.number_format = '0.00'
        c.alignment = Alignment(horizontal='left')
    r += 1
P_CELL = f'README!$B${p_row}'
assert ws.cell(p_row, 1).value.startswith('p (single-pass'), ws.cell(p_row, 1).value
print('p input cell:', P_CELL)

# ---------------------------------------------------------------- data sheets
def data_sheet(title, param_hdr, param_fmt, data, note):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.cell(1, 1, note).font = Font(name=FONT, size=10, italic=True, color='595959')
    ws.cell(2, 1, f"eta_ext and EQE are formulas driven by p in {P_CELL.replace('$','')}").font = \
        Font(name=FONT, size=9, italic=True, color='808080')
    hdr = [param_hdr, 'eta_sub', "A' (round-trip loss)", 'Waveguided', 'u>1 bin',
           'Absorption', 'eta_ext', 'EQE', 'Channel sum (check)']
    for j, h in enumerate(hdr, 1):
        c = ws.cell(4, j, h); c.font = HDR_FONT; c.fill = HDR_FILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BOX
    ws.row_dimensions[4].height = 30
    for i, row in enumerate(data):
        r = 5 + i
        vals = [row[0], row[1], row[2], row[3], row[4], row[5]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, float(v)); c.font = BODY; c.border = BOX
            c.number_format = param_fmt if j == 1 else '0.0000'
        # eta_ext = p / (p + (1-p) A')
        c = ws.cell(r, 7, f'={P_CELL}/({P_CELL}+(1-{P_CELL})*C{r})')
        c.font = BODY; c.border = BOX; c.number_format = '0.0000'
        c = ws.cell(r, 8, f'=B{r}*G{r}')
        c.font = BODY; c.border = BOX; c.number_format = '0.0000'
        c = ws.cell(r, 9, f'=B{r}+D{r}+E{r}+F{r}')
        c.font = Font(name=FONT, size=10, color='808080'); c.border = BOX; c.number_format = '0.000000'
    ws.column_dimensions['A'].width = 22
    for col in 'BCDEFGHI': ws.column_dimensions[col].width = 13
    ws.freeze_panes = 'A5'
    return ws

data_sheet('a_TCO_electrode', 'k_TCO (extinction)', '0.0000', A,
           'Panel (a) — device A: bottom electrode is a 50 nm TCO with n = 1.9 + k_TCO i. Ag reflector fixed at McPeak n,k.')
data_sheet('b_thin_Ag_electrode', 'n_Ag (real index)', '0.00', B,
           'Panel (b) — device B: bottom electrode is a 10 nm Ag with n = n_Ag + 3.5i. Ag reflector fixed at McPeak n,k.')

# ---------------------------------------------------------------- panel c
ws = wb.create_sheet('c_p_sensitivity')
ws.sheet_view.showGridLines = False
ws.cell(1, 1, 'Panel (c) — how the extraction efficiency responds to the microlens escape probability p.').font = \
    Font(name=FONT, size=10, italic=True, color='595959')
ws.cell(2, 1, "Every cell is a formula: eta_ext = p / [p + (1 - p) A'], with A' taken from row 4.").font = \
    Font(name=FONT, size=9, italic=True, color='808080')
cases = [('ideal electrode', A[0, 2]), ('TCO, k_TCO = 0.02', A[8, 2]),
         ('Ag, n_Ag = 0.24', B[12, 2]), ('TCO, k_TCO = 0.08', A[-1, 2])]
c = ws.cell(4, 1, "A'  ->"); c.font = Font(name=FONT, size=10, bold=True, color='808080')
c.alignment = Alignment(horizontal='right')
for j, (lab, Ap) in enumerate(cases, 2):
    c = ws.cell(4, j, float(Ap)); c.font = Font(name=FONT, size=10, bold=True, color='0000FF')
    c.number_format = '0.0000'; c.border = BOX
    c.alignment = Alignment(horizontal='center')
hdr = ['p'] + [lab for lab, _ in cases]
for j, h in enumerate(hdr, 1):
    c = ws.cell(5, j, h); c.font = HDR_FONT; c.fill = HDR_FILL
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True); c.border = BOX
ws.row_dimensions[5].height = 30
for i, p in enumerate(np.round(np.arange(0.15, 0.9001, 0.025), 4)):
    r = 6 + i
    c = ws.cell(r, 1, float(p)); c.font = BODY; c.number_format = '0.000'; c.border = BOX
    for j in range(2, 6):
        col = get_column_letter(j)
        c = ws.cell(r, j, f'=$A{r}/($A{r}+(1-$A{r})*{col}$4)')
        c.font = BODY; c.number_format = '0.0000'; c.border = BOX
ws.column_dimensions['A'].width = 10
for col in 'BCDE': ws.column_dimensions[col].width = 19
ws.freeze_panes = 'A6'

wb.save('design_rule_parasitic_absorption.xlsx')
print('written:', [s.title for s in wb.worksheets])
