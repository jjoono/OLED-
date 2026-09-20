# -*- coding: utf-8 -*-
"""The measured-electrode points of Fig. 3(a) panel (c), as a workbook."""
import json, os, openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
HERE = os.path.dirname(os.path.abspath(__file__))
BOLD = Font(bold=True); YEL = PatternFill('solid', fgColor='FFFF00')
F = json.load(open(os.path.join(HERE, 'films.json')))
wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'README'
ws.column_dimensions['A'].width = 28; ws.column_dimensions['B'].width = 104
rows = [
 ('Fig. 3(a) panel (c) — every electrode in the library on one stack', None),
 ('Each row is a film measured by ellipsometry in this group, taken at 550 nm and run on the same device', None),
 (None, None), ('Stack', None),
 ('Common part', 'Ag 100 nm (McPeak, fixed) / ETL 200 nm / EML 20 nm / HTL 200 nm / [this electrode] / substrate n = 1.80'),
 ('TCO films', '50 nm thick, the thinnest layer that still carries the current'),
 ('thin Ag films', '10 nm thick'),
 ('Settings', '550 nm, isotropic dipole, PLQY = 1, organics isotropic n = 1.8 with k = 0, u grid 3000, p = 0.30'),
 ('Optical constants', 'read from nk_JH_total.mat at 550 nm; the film names are the entry names in that file'),
 (None, None), ('Columns', None),
 ('n, k', 'the measured optical constants of the electrode at 550 nm'),
 ('eta_sub', 'substrate-delivered power'),
 ("A'", 'round-trip loss'),
 ('absorbed in the electrode', 'the part of the dipole power dissipated in this electrode alone — the x axis of the panel'),
 ('absorbed in the reflector', 'the same for the 100 nm Ag mirror, which is common to every row'),
 ('eta_ext, EQE', 'formulas driven by the p cell below'),
 (None, None), ('Input', None), ('p', 0.30),
 (None, None), ('Provenance', None),
 ('Script', 'sim/design_rule4/dr4e.m, one run per film; collected by the shell loop recorded in the README'),
 ('Figure', 'sim/design_rule4/plot_dr4e.py, panel (c); standalone version plot_fig3a_c.py'),
 ('Repository', 'jjoono/OLED-, branch claude/oled-efficiency-paper-ag-955zsk'),
]
for i,(a,b) in enumerate(rows, start=1):
    if a is not None: ws.cell(row=i, column=1, value=a).font = BOLD
    if b is not None:
        c = ws.cell(row=i, column=2, value=b)
        if isinstance(b, float): c.fill = YEL; c.font = BOLD; c.number_format = '0.00'
ws['A1'].font = Font(bold=True, size=14)
P_ROW = [i for i,(a,_) in enumerate(rows, start=1) if a == 'p'][0]

ws = wb.create_sheet('films')
hdr = ['film (library entry)', 'family', 'n at 550 nm', 'k at 550 nm', 'eta_sub', "A'",
       'waveguided', 'u>1', 'absorbed, total', 'absorbed in the reflector',
       'absorbed in the electrode', 'eta_ext', 'EQE']
for j,h in enumerate(hdr, start=1): ws.cell(row=1, column=j, value=h).font = BOLD
P = f'README!$B${P_ROW}'
order = sorted(F, key=lambda f: (f['family'] != 'TCO', f['abs_electrode']))
for i,f in enumerate(order):
    r = 2 + i
    vals = [f['film'], 'TCO 50 nm' if f['family']=='TCO' else 'thin Ag 10 nm', f['n'], f['k'],
            f['eta_sub'], f['Aprime'], f['wg'], f['spp'], f['abs_total'], f['abs_mirror'], f['abs_electrode']]
    for j,v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=j, value=v)
        if isinstance(v, float): c.number_format = '0.0000'
    ws.cell(row=r, column=12, value=f'={P}/({P}+(1-{P})*F{r})').number_format = '0.0000'
    ws.cell(row=r, column=13, value=f'=E{r}*L{r}').number_format = '0.0000'
ws.column_dimensions['A'].width = 20; ws.column_dimensions['B'].width = 15
for j in range(3, 14): ws.column_dimensions[get_column_letter(j)].width = 22 if j in (10, 11) else 14
ws.freeze_panes = 'A2'
wb.save(os.path.join(HERE, 'fig3a_films_rawdata.xlsx'))
print('saved fig3a_films_rawdata.xlsx')
