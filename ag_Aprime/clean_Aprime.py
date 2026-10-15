"""Remove the undersampled Ag/air surface-plasmon spike (theta_sub ~ 43-46 deg) from A'(theta, lambda).

The spike is light coupled through the 100 nm Ag cathode to the plasmon on its air-side face
(k_spp = sqrt(eps/(eps+1)) -> 45.8 deg at 450 nm, 43.0 deg at 700 nm for n_sub = 1.5). It is far
narrower than the 1-degree grid, so it appears as one isolated sample per wavelength. In a real
device the Ag is covered (n >= 1.5) and this plasmon is not reachable from the substrate.
Only points inside 41-48 deg whose deviation from a local cubic fit exceeds THR are replaced."""
import numpy as np, matplotlib
matplotlib.use('Agg'); import matplotlib.pyplot as plt
from openpyxl import Workbook

THR = 0.15          # % points
L = open('Ag_Aprime_raw.tsv').read().strip().split('\n')
th = np.array(L[0].split('\t'), float)
rows = [list(map(float, l.split('\t'))) for l in L[1:]]
lam = np.array([r[0] for r in rows]); A = np.array([r[1:] for r in rows]); C = A.copy()
log = []
W = np.cos(np.radians(th + 0.0)) * np.sin(np.radians(th + 0.0))


def resid(row, c, mask):
    x = np.r_[c - 5:c, c + 1:c + 6]; x = x[mask[x]]
    p = np.polyfit(x, row[x], 3)
    return row[c] - np.polyval(p, c), np.polyval(p, c)


for i in range(len(lam)):
    mask = np.ones(90, bool)
    for _ in range(3):                      # spike + possible Fano partner(s)
        best = max(range(41, 49), key=lambda c: abs(resid(C[i], c, mask)[0]) if mask[c] else -1)
        r, fit = resid(C[i], best, mask)
        if abs(r) < THR:
            break
        mask[best] = False
    bad = np.flatnonzero(~mask)
    if bad.size:
        good = np.r_[35:55]; good = good[mask[good]]
        p = np.polyfit(good, C[i, good], 4)
        for c in bad:
            new = float(np.polyval(p, c)); log.append((lam[i], th[c], C[i, c], new)); C[i, c] = new

Al_raw = (A * W).sum(1) / W.sum(); Al_cln = (C * W).sum(1) / W.sum()
print(f'replaced {len(log)} points in {len(set(l[0] for l in log))} of {len(lam)} wavelengths; '
      f'angles {sorted(set(int(l[1]) for l in log))}')
print('Lambertian-weighted A\' change: max %.4f %%p' % np.abs(Al_raw - Al_cln).max())

# outputs: TSV (same layout), Excel (cleaned, raw, replaced-points log), before/after figure
with open('Ag_Aprime_clean.tsv', 'w') as f:
    f.write('\t'.join(str(int(t)) for t in th) + '\n')
    for i in range(len(lam)):
        f.write(str(int(lam[i])) + '\t' + '\t'.join(f'{v:.5f}' for v in C[i]) + '\n')
wb = Workbook(); ws = wb.active; ws.title = 'A_prime_clean'
ws.append(['lambda_nm \\ theta_sub_deg', *[int(t) for t in th]])
for i in range(len(lam)): ws.append([int(lam[i]), *[round(float(v), 5) for v in C[i]]])
ws = wb.create_sheet('A_prime_raw'); ws.append(['lambda_nm \\ theta_sub_deg', *[int(t) for t in th]])
for i in range(len(lam)): ws.append([int(lam[i]), *[float(v) for v in A[i]]])
ws = wb.create_sheet('replaced_points'); ws.append(['lambda_nm', 'theta_deg', 'raw_%', 'replaced_%'])
for l in log: ws.append([int(l[0]), int(l[1]), float(l[2]), round(l[3], 5)])
ws = wb.create_sheet('Lambertian_avg'); ws.append(['lambda_nm', "A'_raw_%", "A'_clean_%"])
for i in range(len(lam)): ws.append([int(lam[i]), float(Al_raw[i]), float(Al_cln[i])])
wb.save('Ag_Aprime_clean.xlsx')

fig, ax = plt.subplots(1, 3, figsize=(9, 3), gridspec_kw=dict(wspace=0.55))
for k, (M, t) in enumerate(((A, 'raw'), (C, 'cleaned'))):
    im = ax[k].pcolormesh(th, lam, M, cmap='inferno_r', vmin=0, vmax=30, shading='nearest')
    ax[k].set_title(f"(a{k + 1}) A' (%) {t}", loc='left', fontsize=8); ax[k].set_xlabel(r'$\theta$ in substrate (deg)')
    ax[k].set_ylabel('wavelength (nm)')
fig.colorbar(im, ax=ax[:2], fraction=0.03)
for lv, c in ((450, '#3B6EA5'), (520, '#2E7D5B'), (680, '#C6512F')):
    i = int(np.argmin(abs(lam - lv)))
    ax[2].plot(th, A[i], ':', color=c, lw=1); ax[2].plot(th, C[i], '-', color=c, lw=1.2, label=f'{lv} nm')
ax[2].set_xlim(30, 60); ax[2].set_ylim(2, 12); ax[2].set_xlabel(r'$\theta$ in substrate (deg)'); ax[2].set_ylabel("A' (%)")
ax[2].set_title('(b) dotted raw, solid cleaned', loc='left', fontsize=8); ax[2].legend(fontsize=7)
fig.savefig('Ag_Aprime_clean.png', dpi=250, bbox_inches='tight')
