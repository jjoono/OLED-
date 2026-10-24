"""Origin-ready workbook for the three descriptor charts.

    python slides/make_origin_xlsx.py  ->  slides/seed_descriptors_origin.xlsx

Rows 1-3 of each data sheet are Origin's Long Name / Units / Comments; data
start on row 4, sorted ascending so the largest value sits on top of a
horizontal bar chart. Source: runs/slide_descriptors.csv.
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "..", "runs", "slide_descriptors.csv")
OUT = os.path.join(HERE, "seed_descriptors_origin.xlsx")

rows = list(csv.DictReader(open(SRC)))
for r in rows:
    for k in ["E_b", "E_elec", "E_disp", "q_Ag1", "dq_Ag1_2"]:
        r[k] = float(r[k])
    r["label"] = {"Cs2CO3": "Cs₂CO₃"}.get(r["name"], r["name"])


def klass(r):
    if r["name"].startswith("MoOx"):
        return "ref"
    if r["role"].startswith("HIL"):
        return "HIL"
    if r["role"] in ("host", "HTL"):
        return "HOST"
    return "ETL"


org = [r for r in rows if not r["name"].startswith("MoOx")]
moox = next(r for r in rows if r["name"].startswith("MoOx"))

F = Font(name="Arial", size=10)
B = Font(name="Arial", size=10, bold=True)
G = Font(name="Arial", size=10, italic=True, color="666666")
HDR = PatternFill("solid", fgColor="E8EEF7")
wb = Workbook()


def sheet(title, cols, units, comments, data, note):
    ws = wb.create_sheet(title)
    for hdr in (cols, units, comments):
        ws.append(hdr)
    for c in range(1, len(cols) + 1):
        for rr in (1, 2, 3):
            cell = ws.cell(rr, c)
            cell.font = B if rr == 1 else G
            cell.fill = HDR
    for d in data:
        ws.append(d)
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.font = F
            if isinstance(cell.value, float):
                cell.number_format = "0.000"
    ws.column_dimensions["A"].width = 16
    for c in "BCDEFG":
        ws.column_dimensions[c].width = 15
    ws.cell(1, 1).comment = Comment(note, "calc")
    ws.freeze_panes = "B4"
    return ws


d = sorted(org, key=lambda r: r["E_b"])
sheet("Eb", ["Material", "E_elec", "E_disp", "E_b", "Label"],
      ["", "eV", "eV", "eV", ""],
      ["", "화학 결합(전자적)", "분산력", "합계 (PBE0-D3)", "축 라벨: 이름 + E_b"],
      [[r["label"], r["E_elec"], r["E_disp"], r["E_b"], f"{r['label']}  {r['E_b']:.2f}"]
       for r in d],
      "Stacked bar: E_elec + E_disp = E_b. Reference line at 0.5 eV. NPB E_elec = -0.004 (about 0).")

cls = [("HIL", "HIL"), ("ETL", "ETL · EIL"), ("HOST", "호스트 · HTL")]
d = sorted(org, key=lambda r: r["q_Ag1"])
sheet("qAg1", ["Material"] + [c[1] for c in cls] + ["q_Ag1 (all)"],
      [""] + ["e"] * 4, [""] + ["PBE0 Mulliken"] * 4,
      [[r["label"]] + [r["q_Ag1"] if klass(r) == k else None for k, _ in cls] + [r["q_Ag1"]]
       for r in d],
      "One column per class (blank = other class), so each class gets its own colour.")

cls2 = cls + [("ref", "MoOx (참고)")]
d = sorted(org + [moox], key=lambda r: r["dq_Ag1_2"])
sheet("dq", ["Material"] + [c[1] for c in cls2] + ["dq (all)"],
      [""] + ["e"] * 5, [""] + ["GFN2-xTB"] * 5,
      [[r["label"]] + [r["dq_Ag1_2"] if klass(r) == k else None for k, _ in cls2]
       + [r["dq_Ag1_2"]] for r in d],
      "dq(Ag1->Ag2) = total Ag charge with two atoms minus with one. Reference line at +0.3 e.")

ws = sheet("All", ["Material", "Role", "E_b", "E_elec", "E_disp", "q_Ag1", "dq_Ag1_2"],
           ["", "", "eV", "eV", "eV", "e", "e"], ["", "", "PBE0", "PBE0", "D3BJ", "PBE0", "GFN2"],
           [[r["label"], r["role"], r["E_b"], r["E_elec"], r["E_disp"], r["q_Ag1"], r["dq_Ag1_2"]]
            for r in sorted(rows, key=lambda r: -r["E_b"])],
           "Source: runs/slide_descriptors.csv. MoOx: Mo3O9 cluster (E_b, q), MoO3(010) slab (dq).")
ws.column_dimensions["B"].width = 17

rd = wb["Sheet"]
rd.title = "README"
lines = [
    ("Origin에서 불러오기", True),
    ("File > Import > Excel (또는 파일을 Origin 창에 끌어다 놓기). 가져오기 옵션: Long Name = 1행, Units = 2행, Comments = 3행.", False),
    ("시트 하나가 그래프 하나입니다. 가로 막대에서 큰 값이 위로 가도록 오름차순으로 정렬해 두었습니다.", False),
    ("", False),
    ("그래프 만드는 법", True),
    ("Eb: E열(Label)을 X로 지정 → E_elec, E_disp 선택 → Plot > Bar, Pie, Area > Stacked Bar. 색: E_elec #1C5CAB, E_disp #86B6EF. 0.5 eV 기준선 (#D03B3B, 점선).", False),
    ("qAg1: A열을 X로 지정 → 클래스 열 3개 선택 → Plot > Bar, Pie, Area > Stacked Bar (빈 칸은 막대 없음). 색: HIL #EB6834, ETL·EIL #1BAF7A, 호스트·HTL #6250D6. 값 표시: Plot Details > Label 탭, 형식 +0.00.", False),
    ("dq: qAg1과 같고 MoOx 열(#D03B3B) 추가. +0.3 e 탈락선 (#D03B3B, 점선).", False),
    ("저장: File > Save Project As → .opju", False),
    ("", False),
    ("데이터", True),
    ("E_b, E_elec, q_Ag1: PBE0-D3BJ/def2-SVP (큰 분자는 조각 PBE0, 수소에 막히는 자리는 PBE-D3로 Ag 위치를 다시 잡음). E_disp: D3BJ. dq: GFN2-xTB.", False),
    ("CBP, TCTA, TPD, TAPC는 빠져 있습니다. mCP, NPB와 같은 π 계열입니다 (0.25–0.28 eV).", False),
    ("값은 수식이 아니라 계산 결과입니다. 계산이 바뀔 때만 고치세요.", False),
]
for i, (t, b) in enumerate(lines, 1):
    c = rd.cell(i, 1, t)
    c.font = B if b else F
    c.alignment = Alignment(wrap_text=True, vertical="top")
rd.column_dimensions["A"].width = 120
wb.save(OUT)
print("saved", OUT, wb.sheetnames)
