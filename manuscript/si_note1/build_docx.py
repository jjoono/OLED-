import docx, sys
from lxml import etree
from eqs import EQ
d = docx.Document()
d.add_heading('Supplementary Note 1 — 수식', 1)
for n, title, body, note in EQ:
    d.add_paragraph().add_run(f'({n}) {title}').bold = True
    p = d.add_paragraph(); p._p.append(etree.fromstring(__import__('omml').para(body)))
    d.add_paragraph(note)
d.save('note1_equations.docx')
