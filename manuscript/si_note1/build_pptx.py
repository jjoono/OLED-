import copy, re, csv, numpy as np
from lxml import etree
from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.dml.color import RGBColor
import omml
from eqs import EQ
A14 = 'http://schemas.microsoft.com/office/drawing/2010/main'
MC = 'http://schemas.openxmlformats.org/markup-compatibility/2006'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'
AN = 'http://schemas.openxmlformats.org/drawingml/2006/main'
RPR = '<a:rPr xmlns:a="%s" lang="en-US" sz="2000"><a:latin typeface="Cambria Math" panose="02040503050406030204"/><a:cs typeface="Cambria Math"/></a:rPr>' % AN
def mathify(xml):   # PowerPoint math runs carry a DrawingML run property with the math font
    return re.sub(r'<m:r>(<m:rPr>.*?</m:rPr>)?', lambda m: '<m:r>' + (m.group(1) or '') + RPR, xml)
prs = Presentation(); prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)
blank = prs.slide_layouts[6]
def text(s, x, y, w, h, txt, size=14, bold=False, color=None):
    tb = s.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h)); tf = tb.text_frame; tf.word_wrap = True
    for i, line in enumerate(txt.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); r = p.add_run(); r.text = line
        r.font.size = Pt(size); r.font.bold = bold
        if color: r.font.color.rgb = RGBColor(*color)
    return tb
def equation(s, x, y, w, h, body, fallback):
    tb = text(s, x, y, w, h, fallback, 20)
    sp = tb._element; fb = copy.deepcopy(sp)
    p = sp.find('.//{%s}p' % AN)
    for ch in list(p): p.remove(ch)
    m = etree.SubElement(p, '{%s}m' % A14)
    m.append(etree.fromstring(mathify('<m:oMathPara xmlns:m="%s"><m:oMathParaPr><m:jc m:val="left"/></m:oMathParaPr><m:oMath>%s</m:oMath></m:oMathPara>' % (omml.M, body))))
    ac = etree.Element('{%s}AlternateContent' % MC, nsmap={'mc': MC})
    ch = etree.SubElement(ac, '{%s}Choice' % MC, nsmap={'a14': A14}); ch.set('Requires', 'a14')
    parent = sp.getparent(); idx = parent.index(sp); parent.remove(sp); ch.append(sp)
    etree.SubElement(ac, '{%s}Fallback' % MC).append(fb); parent.insert(idx, ac)
# ---- title
s = prs.slides.add_slide(blank)
text(s, 2, 5, 30, 3, 'Supplementary Note 1 — 수식 (편집용)', 32, True)
text(s, 2, 8.5, 30, 6, '본문 기호와 통일: B_T, B_R, R_LED, p, A′, η_sub, η_ext\n'
     '통과 번호 k(= 학위논문의 N_R): P_sub^(k)(θ)는 k번째로 광추출 구조에 도달하는 빛(정규화하지 않음)\n'
     '새로 도입: 통과별 탈출 확률 p_k, 통과별 왕복 손실 A′_k → 식 (S6)–(S8), Supplementary Fig. (p_k, A′_k)\n'
     '모든 수식은 PowerPoint 수식 개체입니다(더블클릭으로 편집).', 16)
# ---- equations, two per slide
for i in range(0, len(EQ), 2):
    s = prs.slides.add_slide(blank); y = 0.8
    for n, title, body, note in EQ[i:i + 2]:
        text(s, 1.2, y, 31, 1.2, f'({n})  {title}', 20, True, (0x1f, 0x3a, 0x68))
        equation(s, 1.8, y + 1.4, 30, 3.2, body, f'({n}) see Word version')
        text(s, 1.8, y + 4.9, 30, 2.4, note, 14, color=(0x40, 0x40, 0x40))
        y += 9.0
# ---- mapping slide
s = prs.slides.add_slide(blank)
text(s, 1.2, 0.8, 31, 1.2, '학위논문(슬라이드) 표기 → Note 1 표기', 22, True, (0x1f, 0x3a, 0x68))
rows = [('학위논문 / 슬라이드', 'Note 1', '비고'),
        ('BSDF_T, BSDF_R', 'B_T, B_R', '본문 Methods와 통일'),
        ('N_R = k, (R = k)', '(k)', '한 표기로 통일'),
        ('η_sub^(N_R=k) × P_sub^(N_R=k) (k ≥ 1)', 'P_sub^(k) (정규화하지 않음)', 'η_sub는 k = 0(소자→기판)에만'),
        ('η_out^(N_R=k)', 'η_ext^(k)', 'η_sub로 나눈 통과별 추출량'),
        ('η_out^Total', 'η_out = η_sub · η_ext', '본문 식 (1), (2)'),
        ('P̄_sub × BSDF_T', 'p (또는 p_k)', '탈출 확률'),
        ('A′_para', 'A′ (또는 A′_k)', 'OLED에서는 A′_act = 0'),
        ('A′_act, η*_rad, m', '—', 'PeLED 재발광 항: OLED에서 0, 문장으로만 언급'),
        ('BSDF_R ≈ 1 − BSDF_T', 'Σ_θr B_R = 1 − B_T (등호)', '흡수 없는 광추출 구조에서 정확')]
tbl = s.shapes.add_table(len(rows), 3, Cm(1.5), Cm(2.6), Cm(30.5), Cm(14)).table
for i, rw in enumerate(rows):
    for j, v in enumerate(rw):
        c = tbl.cell(i, j); c.text = v
        for p in c.text_frame.paragraphs:
            for r in p.runs: r.font.size = Pt(14); r.font.bold = (i == 0)
# ---- SI figure slide
s = prs.slides.add_slide(blank)
text(s, 1.2, 0.6, 31, 1.2, 'Supplementary Fig. (제안) — 통과별 p_k, A′_k와 누적 η_ext', 22, True, (0x1f, 0x3a, 0x68))
s.shapes.add_picture('pass_resolved_preview.png', Cm(1.2), Cm(2.2), width=Cm(22))
text(s, 23.8, 2.4, 9.5, 15,
     '그림 2(d)–(f) 스택, 550 nm, 반구 MLA(n_MLA = n_sub).\n\n'
     '(a) p_k: 통과를 거듭할수록 감소. 되돌아온 빛이 큰 각도로 몰리기 때문(60° 이상 비율: Ag 1.5에서 29 → 51 %, k = 0 → 2).\n\n'
     '(b) A′_k: 같은 이유로 증가(Ag 1.5: 3.4 → 4.6 %, k = 0 → 5).\n\n'
     '(c) 누적 η_ext: 급수(실선)가 식 (3)(점선)보다 낮게 수렴. Ag 1.5: 0.911 대 0.964; Al 1.5: 0.782 대 0.840.\n\n'
     '가로 점선: 식 (3)에 쓴 p(cos θ sin θ)와 A′.\n데이터: sim/pass_resolved/pass_resolved.csv', 13)
prs.save('note1_equations.pptx')
print('saved', len(prs.slides._sldIdLst), 'slides')
