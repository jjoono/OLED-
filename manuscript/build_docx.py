# -*- coding: utf-8 -*-
"""Build unityEQE_v18_ko.docx from unityEQE_final_ko.md, using the author's uploaded
docx as the style template so fonts and page setup stay theirs.

    python3 build_docx.py <template.docx> <in.md> <out.docx>

Markdown subset: '#'..'####' headings (bold runs, the top one larger), '**...**' bold
spans, '- ' items (kept as plain paragraphs with an en dash, no literal bullets),
blank lines, everything else a Normal paragraph.
"""
import re, sys, copy
import docx
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

template, src, out = sys.argv[1:4]
d = docx.Document(template)
body = d.element.body
# clear the body but keep the final section properties
for el in list(body):
    if el.tag.endswith('}sectPr'):
        continue
    body.remove(el)

BOLD_RE = re.compile(r'\*\*(.+?)\*\*')


def add_runs(p, text, bold_all=False, size=None):
    pos = 0
    for m in BOLD_RE.finditer(text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()]); r.bold = bold_all
            if size: r.font.size = size
        r = p.add_run(m.group(1)); r.bold = True
        if size: r.font.size = size
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:]); r.bold = bold_all
        if size: r.font.size = size


lines = open(src, encoding='utf-8').read().split('\n')
for line in lines:
    s = line.rstrip()
    if not s:
        d.add_paragraph()
        continue
    m = re.match(r'^(#{1,4})\s+(.*)$', s)
    if m:
        level = len(m.group(1))
        p = d.add_paragraph()
        size = {1: Pt(14), 2: Pt(12), 3: Pt(11), 4: Pt(10.5)}[level]
        add_runs(p, m.group(2), bold_all=True, size=size)
        p.paragraph_format.space_before = Pt(8 if level > 1 else 0)
        p.paragraph_format.space_after = Pt(4)
        continue
    if s.startswith('- '):
        p = d.add_paragraph()
        add_runs(p, '– ' + s[2:])
        p.paragraph_format.left_indent = Pt(12)
        continue
    if s.startswith('---'):
        continue
    p = d.add_paragraph()
    # centred, single-line equation paragraphs
    if re.search(r'\s\(\d\)$', s) and ('=' in s) and len(s) < 90:
        add_runs(p, s)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        continue
    add_runs(p, s)

d.save(out)
print('wrote', out, 'paragraphs', len(d.paragraphs))
