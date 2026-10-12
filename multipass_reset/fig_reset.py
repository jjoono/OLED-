"""Draft Fig. S8: angle-preserving rear vs Lambertian-reset rear, common cavity."""
import numpy as np, csv, matplotlib
matplotlib.use('Agg'); import matplotlib.pyplot as plt
import multipass as MP
from openpyxl import Workbook
D = np.load('bsdfs.npz')
L = ['hemi_0-20','hemi_20-40','hemi_40-60','hemi_60-80','hemi_total','free_20-40','free_40-60','free_60-80','free_total']
lab = ['H 0-20','H 20-40','H 40-60','H 60-80','H total','F 20-40','F 40-60','F 60-80','F total']
out = {}
for cat, refl in (('Al', MP.AL), ('Ag', MP.M.AG)):
    es, R, P = MP.device(refl, 61.3, 97.5)
    for d in L:
        BT, BR = D[d+'_BT'], D[d+'_BR']
        r = {}
        for mode in ('spec', 'reset'):
            e, a, pk = MP.run(BT, BR, R, P, mode); r[mode] = (e, MP.comp(a), pk)
        r['dA'] = [MP.run(BT, BR, R, P, 'reset', 1.0, x)[0] for x in np.linspace(0, 0.05, 11)]
        out[(cat, d)] = r
fig, ax = plt.subplots(1, 3, figsize=(7.4, 2.6), gridspec_kw=dict(wspace=0.55))
# (a) p_k
for mode, c, nm in (('spec', '#3B6EA5', 'angle-preserving'), ('reset', '#C6512F', 'Lambertian reset')):
    pk = out[('Al', 'hemi_total')][mode][2][:12]
    ax[0].plot(np.arange(1, 13), pk, 'o-', ms=3, color=c, label=nm)
ax[0].set_xlabel('pass k'); ax[0].set_ylabel('escape probability $p_k$'); ax[0].set_title('(a) hemisphere, Al', loc='left')
ax[0].legend(fontsize=6); ax[0].grid(alpha=.25)
# (b) eta_ext all lenses
x = np.arange(len(L))
for cat, mk in (('Al', 'o'), ('Ag', 's')):
    ax[1].plot(x, [out[(cat, d)]['spec'][0] for d in L], mk, mfc='w', color='#3B6EA5', ms=4, label=f'{cat}, angle-preserving')
    ax[1].plot(x, [out[(cat, d)]['reset'][0] for d in L], mk, color='#C6512F', ms=4, label=f'{cat}, reset')
ax[1].set_xticks(x); ax[1].set_xticklabels(lab, rotation=70, fontsize=5.5); ax[1].set_ylabel(r'$\eta_{\rm sta}$')
ax[1].set_title('(b) 9 lenses, one cavity', loc='left'); ax[1].legend(fontsize=5, ncol=1); ax[1].grid(alpha=.25)
# (c) composition spread across lenses
B = ['0-20', '20-40', '40-60', '60-80']
for k, (mode, c) in enumerate((('spec', '#3B6EA5'), ('reset', '#C6512F'))):
    S = np.array([out[('Al', d)][mode][1] for d in L])
    ax[2].errorbar(np.arange(4) + (k - .5) * .25, S.mean(0), yerr=[S.mean(0) - S.min(0), S.max(0) - S.mean(0)], fmt='o', ms=3, color=c, capsize=2,
                   label='angle-preserving' if mode == 'spec' else 'Lambertian reset')
ax[2].plot(np.arange(4), MP.LAMB, 'k_', ms=10, label='Lambertian')
ax[2].set_xticks(range(4)); ax[2].set_xticklabels(B, fontsize=6); ax[2].set_ylabel('$S_j$ (range over 9 lenses)')
ax[2].set_title('(c) composition (Al)', loc='left'); ax[2].legend(fontsize=5); ax[2].grid(alpha=.25)
fig.savefig('figS8_angle_reset_draft.png', dpi=300, bbox_inches='tight'); fig.savefig('figS8_angle_reset_draft.pdf', bbox_inches='tight')
wb = Workbook(); ws = wb.active; ws.title = 'a_pass_resolved'
ws.append(['pass k', 'p_k angle-preserving', 'p_k Lambertian reset'])
for k in range(12): ws.append([k + 1, float(out[('Al','hemi_total')]['spec'][2][k]), float(out[('Al','hemi_total')]['reset'][2][k])])
ws = wb.create_sheet('b_eta_sta'); ws.append(['lens', 'Al angle-preserving', 'Al reset', 'Ag angle-preserving', 'Ag reset'])
for d in L: ws.append([d, *[float(out[(c, d)][m][0]) for c in ('Al', 'Ag') for m in ('spec', 'reset')]])
ws = wb.create_sheet('c_composition'); ws.append(['lens', 'rear', *['S_' + b for b in B]])
for d in L:
    for m in ('spec', 'reset'): ws.append([d, m, *map(float, out[('Al', d)][m][1])])
ws = wb.create_sheet('d_extra_loss'); ws.append(['extra round-trip loss', *[f'{c} {d}' for c in ('Al','Ag') for d in L]])
for i, x in enumerate(np.linspace(0, 0.05, 11)):
    ws.append([float(x), *[float(out[(c, d)]['dA'][i] / out[(c, d)]['spec'][0]) for c in ('Al','Ag') for d in L]])
wb.save('figS8_angle_reset_draft.xlsx')
for c in ('Al','Ag'):
    g = np.array([out[(c,d)]['dA'] for d in L]) / np.array([out[(c,d)]['spec'][0] for d in L])[:,None]
    xs = np.linspace(0,0.05,11); be = [np.interp(1.0, gi[::-1], xs[::-1]) for gi in g]
    print(c, 'break-even extra loss (pp): min %.1f median %.1f max %.1f' % (100*min(be), 100*np.median(be), 100*max(be)))
